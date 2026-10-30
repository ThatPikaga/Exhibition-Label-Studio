#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
EXHIBITION LABEL STUDIO
=======================
Turns a Microsoft Forms / Excel / CSV export of artwork entries into printable
gallery labels (PDF).  Everything lives in this one file.

    python Generate_Labels.py               make the labels
    python Generate_Labels.py configure     open the friendly set-up screen

Folders (created automatically):
    EXCEL_FILES_HERE/   drop your .xlsx / .csv exports here
    profiles/           one small file per form layout  (e.g. "What Orana Means To Me - 2026.json")
    Generated_Labels/   finished PDFs appear here
"""
import sys, os, re, csv, copy, json, shutil, argparse, importlib, subprocess, textwrap, unicodedata, atexit
from html import unescape as html_unescape
from xml.sax.saxutils import escape as xml_escape
from datetime import datetime


# ════════════════════════════════════════════════════════════════════════════
#  First-run helper: install the two libraries we need if they are missing
# ════════════════════════════════════════════════════════════════════════════
def _ensure_packages():
    needed = []
    for mod, pkg in (("openpyxl", "openpyxl"), ("reportlab", "reportlab")):
        try:
            importlib.import_module(mod)
        except ImportError:
            needed.append(pkg)
    if not needed:
        return
    print("\n  First-time setup: installing %s (needs internet, one time only)...\n" % ", ".join(needed))
    try:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "--quiet"] + needed)
    except Exception:
        print("\n  Could not install the required packages automatically.")
        print("  Please run:  pip install " + " ".join(needed))
        sys.exit(1)
    importlib.invalidate_caches()


_ensure_packages()

from openpyxl import load_workbook
from reportlab.lib import colors
from reportlab.lib.pagesizes import A3, A4, letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph

IS_WIN = os.name == "nt"
if IS_WIN:
    import msvcrt
else:
    import termios, tty, select

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

# ════════════════════════════════════════════════════════════════════════════
#  Paths & settings
# ════════════════════════════════════════════════════════════════════════════
BASE = os.path.dirname(os.path.abspath(__file__))
INPUT_FOLDER = os.path.join(BASE, "EXCEL_FILES_HERE")
OUTPUT_FOLDER = os.path.join(BASE, "Generated_Labels")
PROFILES_DIR = os.path.join(BASE, "profiles")
DATA_EXTS = (".xlsx", ".xlsm", ".csv")

MIN_FONT = 10.0          # statement text is never smaller than this (points)
PAD = 14                 # space inside each label edge (points)
RULE_GAP = 16            # space above the artist statement (points)
PAGE_SIZES = {"A4": A4, "A3": A3, "Letter": letter}

DEFAULT_LAYOUT = dict(page_size="A4", orientation="landscape", columns=2, rows=2, margin_mm=10,
                      statement_size=12, statement_min=10, sort="file")
DEFAULT_AUCTION = dict(yes_words=["yes", "y", "true"], text_yes="Available for Auction", text_no="Not available for Auction")

# (key, short name, question, required, words that hint at the right column)
FIELDS = [
    ("artist", "Artist name", "Which column holds the ARTIST'S NAME?", True,
     ["full name", "artist", "participant", "student", "name"]),
    ("role", "Year level / class", "Which column holds the artist's YEAR LEVEL or CLASS?  (shown under the name)", False,
     ["year level", "year", "class", "grade", "role", "form"]),
    ("title", "Artwork title", "Which column holds the ARTWORK TITLE?", True,
     ["title", "name of your artwork", "artwork"]),
    ("statement", "Artist statement", "Which column holds the ARTIST STATEMENT?  (the longer paragraph)", False,
     ["statement", "describe", "about your", "story"]),
    ("auction", "Auction status", "Which column says whether the artwork is for the SILENT AUCTION?  (a Yes / No answer)", False,
     ["auction", "for sale", "sell"]),
    ("permission", "Display permission", "Which column holds PERMISSION TO DISPLAY?  (optional - you'll get a warning if anyone didn't say Yes)", False, ["display", "permission", "consent"]),
]
FIELD_NAME = {k: n for k, n, *_ in FIELDS}
FIELD_KEYS = tuple(k for k, *_ in FIELDS)

DEFAULT_PROFILE = {
    "profile_name": "What Orana Means To Me - 2026",
    "fields": {
        # These are the clear headings in the cleaned export. resolve() also
        # recognises the longer Microsoft Forms question headings.
        "artist": "Artist",
        "role": "Year / role",
        "title": "Artwork title",
        "statement": "Artist statement",
        # This export doesn't have a confirmed auction-status column, so don't
        # invent a red/green auction marker when the source doesn't say.
        "auction": "",
        "permission": "Display permission",
    },
    "details": [
        {"label": "", "column": "Medium"},
        {"label": "", "column": "Dimensions / size"},
    ],
    "auction": copy.deepcopy(DEFAULT_AUCTION),
    "layout": copy.deepcopy(DEFAULT_LAYOUT),
    "form_columns": [],
}

NA_VALUES = {"na", "n/a", "n.a.", "none", "nil", "-", "--", "not applicable"}


# ════════════════════════════════════════════════════════════════════════════
#  Small text helpers
# ════════════════════════════════════════════════════════════════════════════
def norm(s):
    return re.sub(r"\s+", " ", str(s).replace("\xa0", " ")).strip().casefold()


def header_key(h):
    """A short, stable name for a column header (its first line, lower-case)."""
    lines = [l.strip() for l in str(h).replace("\xa0", " ").splitlines() if l.strip()]
    k = norm(lines[0]) if lines else ""
    k = re.sub(r"(\.\.\.|…)+$", "", k).strip()
    return k[:80].strip()


def header_phrase(h):
    """What we save in a profile for a column: its first line, tidied, cut at a whole word (max 80 characters)."""
    lines = [l.strip() for l in str(h).replace("\xa0", " ").splitlines() if l.strip()]
    t = re.sub(r"\s+", " ", lines[0] if lines else "").strip()
    t = re.sub(r"(\.\.\.|…)+$", "", t).strip()
    if len(t) > 80:
        t = t[:80].rsplit(" ", 1)[0]
    return t


def short(s, n):
    s = re.sub(r"\s+", " ", str(s)).strip()
    return s if len(s) <= n else s[: max(1, n - 1)].rstrip() + "…"


def nice_header(h, n=60):
    lines = [l.strip() for l in str(h).replace("\xa0", " ").splitlines() if l.strip()]
    return short(lines[0] if lines else "(blank column)", n)


def clean_text(val, multi=True):
    """Tidy spreadsheet text and turn form HTML line breaks into real new lines.

    Microsoft Forms / exported listing text can contain literal <br> tags (or
    encoded versions such as &lt;br&gt;). Decode those before PDF markup escapes
    the text, otherwise the tags are printed visibly instead of creating breaks.
    """
    if val is None:
        return ""
    if isinstance(val, float) and val.is_integer():
        val = int(val)
    s = str(val).replace("\xa0", " ").replace("\u200b", "")
    # Decode HTML entities first, so an encoded &lt;br&gt; is handled too.
    s = html_unescape(s)
    # Preserve intentional paragraphs/line breaks from forms and web listings.
    s = re.sub(r"(?i)<br\s*/?\s*>", "\n", s)
    s = re.sub(r"(?i)</p\s*>", "\n\n", s)
    s = re.sub(r"(?i)<p(?:\s+[^>]*)?>", "", s)
    s = re.sub(r"(?i)</div\s*>", "\n", s)
    s = re.sub(r"(?i)<div(?:\s+[^>]*)?>", "", s)
    s = s.replace("\r\n", "\n").replace("\r", "\n").strip()
    if multi and s.endswith(";"):
        s = ", ".join(p.strip() for p in s.split(";") if p.strip())
    return s


def is_yes(text, words):
    t = norm(text)
    return any(re.match(re.escape(norm(w)) + r"\b", t) for w in words if norm(w))


def safe_filename(name):
    return re.sub(r'[\\/:*?"<>|]+', "", str(name)).strip().rstrip(".") or "Untitled"


def stmt_markup(text):
    # Normalise form-supplied <br> tags here as well as in clean_text, so any
    # statement passed to the PDF renderer always turns line breaks into layout.
    text = clean_text(text, multi=False)
    lines = [l.strip() for l in str(text).replace("\r", "").split("\n")]
    t = re.sub(r"\n{3,}", "\n\n", "\n".join(lines)).strip()
    return xml_escape(t).replace("\n", "<br/>")


# ════════════════════════════════════════════════════════════════════════════
#  Reading spreadsheets
# ════════════════════════════════════════════════════════════════════════════
def find_data_files():
    if not os.path.isdir(INPUT_FOLDER):
        return []
    return [os.path.join(INPUT_FOLDER, f) for f in sorted(os.listdir(INPUT_FOLDER), key=str.casefold)
            if f.lower().endswith(DATA_EXTS) and not f.startswith(("~$", "."))]


def read_table(path):
    """Return (headers, rows).  rows = [(sheet_row_number, [cell, ...]), ...]"""
    if path.lower().endswith(".csv"):
        data, last = None, None
        for enc in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                with open(path, newline="", encoding=enc) as fh:
                    data = list(csv.reader(fh))
                break
            except UnicodeDecodeError as e:
                last = e
        if data is None:
            raise last
    else:
        wb = load_workbook(path, data_only=True)   # (read-only mode mis-reads MS Forms files)
        try:
            data = [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
        finally:
            wb.close()
    numbered = [(i + 1, r) for i, r in enumerate(data)]
    while numbered and not any(clean_text(v) for v in numbered[0][1]):
        numbered.pop(0)
    if not numbered:
        return [], []
    headers = ["" if h is None else str(h) for h in numbered[0][1]]
    rows = [(n, r) for n, r in numbered[1:] if any(clean_text(v) for v in r)]
    return headers, rows


def find_col(headers, phrase):
    """Find the column matching a saved phrase: exact > starts-with > contained-in-header."""
    p = norm(phrase)
    if not p:
        return None
    keys = [header_key(h) for h in headers]
    for i, k in enumerate(keys):
        if k == p:
            return i
    for i, k in enumerate(keys):
        if k.startswith(p):
            return i
    for i, h in enumerate(headers):
        if p in norm(h):
            return i
    return None


def guess_column(headers, hints, taken=()):
    """Best-guess column for a field, using words that usually appear in its question."""
    best, best_score = None, 0
    for i, h in enumerate(headers):
        k, full = header_key(h), norm(h)
        if i in taken or not k or k in ("id", "email", "start time", "completion time", "last modified time"):
            continue
        for rank, w in enumerate(hints):
            score = (len(hints) - rank) if w in k else (len(hints) - rank) * 0.5 if w in full else 0
            if score > best_score:
                best, best_score = i, score
    return best


# ════════════════════════════════════════════════════════════════════════════
#  Profiles  (one JSON file per form layout, kept in ./profiles)
# ════════════════════════════════════════════════════════════════════════════
def normalize_profile(p):
    p.setdefault("profile_name", "Untitled")
    p["fields"] = {**{k: "" for k in FIELD_KEYS}, **p.get("fields", {})}
    p.setdefault("details", [])
    p["auction"] = {**DEFAULT_AUCTION, **p.get("auction", {})}
    if not p["auction"].get("text_no"):
        p["auction"]["text_no"] = DEFAULT_AUCTION["text_no"]
    p["layout"] = {**DEFAULT_LAYOUT, **p.get("layout", {})}
    p.setdefault("form_columns", [])
    return p


def new_profile(name):
    return normalize_profile({"profile_name": name})


def profile_path(name):
    return os.path.join(PROFILES_DIR, safe_filename(name) + ".json")


def save_profile(p):
    os.makedirs(PROFILES_DIR, exist_ok=True)
    path = profile_path(p["profile_name"])
    old = p.get("_file")
    data = {k: v for k, v in p.items() if not k.startswith("_")}
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(data, fh, indent=2, ensure_ascii=False)
    if old and os.path.abspath(old) != os.path.abspath(path) and os.path.exists(old):
        os.remove(old)
    p["_file"] = path


BAD_PROFILES = []


def load_profiles():
    """Load every profile.  Creates the folder (with the Orana example) the first time."""
    fresh = not os.path.isdir(PROFILES_DIR)
    os.makedirs(PROFILES_DIR, exist_ok=True)
    if fresh:
        save_profile(normalize_profile(copy.deepcopy(DEFAULT_PROFILE)))
    out, BAD_PROFILES[:] = [], []
    for fn in sorted(os.listdir(PROFILES_DIR), key=str.casefold):
        if not fn.lower().endswith(".json"):
            continue
        path = os.path.join(PROFILES_DIR, fn)
        try:
            with open(path, encoding="utf-8") as fh:
                p = normalize_profile(json.load(fh))
            p["_file"] = path
            out.append(p)
        except Exception:
            BAD_PROFILES.append(fn)
    return out


# Alternate headings used by the cleaned workbook and common Microsoft Forms exports.
# Exact matches are preferred so, for example, "Artist" never accidentally maps
# to the "Artist statement" column.
_NAME_PERMISSION_ALIASES = [
    "i give permission to display my work with my name",
    "permission to display my work with my name",
    "display my work anonymously",
    "without my name",
]
_HEADER_ALIASES = {
    "artist": ["artist", "artist name", "participant name", "student name", "your full name (participant)", "full name"],
    "role": ["year / role", "year / class", "year level / class", "year level", "year / class (shown under the name)", "class", "role"],
    "title": ["artwork title", "title of your artwork", "name of your artwork", "title"],
    "statement": ["artist statement", "artist's statement", "artist statement (the longer paragraph)", "artwork description", "statement"],
    "auction": ["auction status", "included in the silent auction", "silent auction status", "auction", "for sale"],
    "permission": ["display permission", "permission to display with name", "permission to display my work with my name", "i give permission to display my work with my name", "display consent", "consent", "permission"],
}
_DETAIL_HEADER_ALIASES = {
    "medium": ["medium", "mediums", "mediums used", "materials", "what mediums did you use"],
    "dimension": ["dimensions / size", "dimensions", "size", "what are the dimensions of your physical piece"],
}


def _compact_header(value):
    return re.sub(r"[^a-z0-9]+", " ", norm(value)).strip()


def _find_alias_col(headers, aliases, taken=()):
    """Match alternate headings, preferring exact names over partial matches."""
    available = [(i, header_key(h), _compact_header(h)) for i, h in enumerate(headers) if i not in taken]
    alias_keys = [_compact_header(a) for a in aliases]
    # Exact match first.
    for alias in alias_keys:
        for i, key, full in available:
            if key and (key == alias or full == alias):
                return i
    # Then match a heading that begins with a known phrase.
    for alias in alias_keys:
        if len(alias) < 5:
            continue
        for i, key, full in available:
            if key.startswith(alias) or full.startswith(alias + " "):
                return i
    # Last resort: a longer, distinctive alias contained in the heading.
    for alias in sorted((a for a in alias_keys if len(a) >= 9), key=len, reverse=True):
        for i, key, full in available:
            if alias in full:
                return i
    return None


def resolve(profile, headers):
    """Resolve saved profile columns, with aliases for changed export headings."""
    idx = {}
    used = set()
    for key in FIELD_KEYS:
        phrase = profile["fields"].get(key, "")
        column = find_col(headers, phrase) if phrase else None
        if column in used:
            column = None
        # Prefer a specific name-display consent question over an older generic
        # "artwork may be displayed" question when both exist in the form export.
        if key == "permission":
            specific_permission = _find_alias_col(headers, _NAME_PERMISSION_ALIASES, used)
            if specific_permission is not None:
                column = specific_permission
        if column is None:
            column = _find_alias_col(headers, _HEADER_ALIASES.get(key, []), used)
        idx[key] = column
        if column is not None:
            used.add(column)

    det = []
    for detail in profile["details"]:
        label = detail.get("label", "")
        phrase = detail.get("column", "")
        column = find_col(headers, phrase) if phrase else None
        if column in used:
            column = None
        if column is None:
            detail_key = "medium" if any(x in norm(label + " " + phrase) for x in ("medium", "material")) else \
                         "dimension" if any(x in norm(label + " " + phrase) for x in ("dimension", "size")) else ""
            column = _find_alias_col(headers, _DETAIL_HEADER_ALIASES.get(detail_key, []), used) if detail_key else None
        det.append((label, column, phrase))
        if column is not None:
            used.add(column)
    return idx, det


def match_score(profile, headers):
    """0..1 - how many of this profile's configured columns can be resolved."""
    idx, details = resolve(profile, headers)
    want = [key for key, value in profile["fields"].items() if value]
    want += ["artist", "title"] if not want else []
    detail_want = [d for d in profile["details"] if d.get("column")]
    total = len(want) + len(detail_want)
    if total == 0:
        return 0.0
    matched = sum(idx.get(key) is not None for key in want)
    matched += sum(col is not None for label, col, phrase in details if phrase)
    return matched / total


# ════════════════════════════════════════════════════════════════════════════
#  Fonts  (uses a real Unicode font when one is installed, so accents & symbols work)
# ════════════════════════════════════════════════════════════════════════════
FONTS = {"reg": "Helvetica", "bold": "Helvetica-Bold", "ital": "Helvetica-Oblique", "cmap": None, "loaded": False}
_FAMILIES = [
    ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf"),
    ("liberationsans-regular.ttf", "liberationsans-bold.ttf", "liberationsans-italic.ttf", "liberationsans-bolditalic.ttf"),
    ("notosans-regular.ttf", "notosans-bold.ttf", "notosans-italic.ttf", "notosans-bolditalic.ttf"),
    ("dejavusans.ttf", "dejavusans-bold.ttf", "dejavusans-oblique.ttf", "dejavusans-boldoblique.ttf"),
    ("segoeui.ttf", "segoeuib.ttf", "segoeuii.ttf", "segoeuiz.ttf"),
]
_FONT_DIRS = [os.environ.get("WINDIR", r"C:\Windows") + r"\Fonts", "/usr/share/fonts", "/usr/local/share/fonts",
              os.path.expanduser("~/.fonts"), os.path.expanduser("~/.local/share/fonts"),
              "/Library/Fonts", "/System/Library/Fonts/Supplemental"]


def load_fonts():
    if FONTS["loaded"]:
        return
    FONTS["loaded"] = True
    found = {}
    for d in _FONT_DIRS:
        for root, _, files in os.walk(d):
            for f in files:
                if f.lower().endswith(".ttf"):
                    found.setdefault(f.lower(), os.path.join(root, f))
    for fam in _FAMILIES:
        if all(f in found for f in fam):
            try:
                for name, f in zip(("LabelReg", "LabelBold", "LabelItal", "LabelBoldItal"), fam):
                    pdfmetrics.registerFont(TTFont(name, found[f]))
                FONTS.update(reg="LabelReg", bold="LabelBold", ital="LabelItal",
                             cmap=pdfmetrics.getFont("LabelReg").face.charToGlyph)
                return
            except Exception:
                continue


def unsupported(text):
    """Characters the label font cannot draw."""
    bad = []
    for ch in str(text):
        if ch in bad or unicodedata.category(ch) in ("Zs", "Cf", "Cc", "Zl", "Zp"):
            continue
        if FONTS["cmap"] is not None:
            ok = ord(ch) in FONTS["cmap"]
        else:
            try:
                ch.encode("cp1252")
                ok = True
            except UnicodeEncodeError:
                ok = False
        if not ok:
            bad.append(ch)
    return bad


# ════════════════════════════════════════════════════════════════════════════
#  Turning spreadsheet rows into label "items"
# ════════════════════════════════════════════════════════════════════════════
def _hide_artist_name(permission_text):
    """Return True when the response explicitly asks for the artist's name to be hidden."""
    answer = norm(permission_text)
    if not answer:
        return False
    if "anonymously" in answer or "anonymous" in answer or "without my name" in answer:
        return True
    # A direct No to displaying the work with the artist's name means the label
    # should not print the artist's name. It doesn't automatically mean the art
    # itself cannot be displayed.
    if answer.startswith("no") and "name" in answer and ("display" in answer or "permission" in answer):
        return True
    return False


def _prepare_statement_for_label(statement, where, who, notes):
    """Remove our own truncation markers without printing them on a gallery label.

    We cannot recreate text that wasn't supplied. If a truncation marker exists,
    keep only the complete sentences before the unfinished sentence and warn the
    user to check the original response.
    """
    text = clean_text(statement, multi=False)
    marker = re.search(
        r"\[\s*(?:the supplied text ends here|please check the original response|statement incomplete in source)[^\]]*\]",
        text, flags=re.IGNORECASE,
    )
    if not marker:
        return text

    complete = text[:marker.start()].rstrip()
    sentence_ends = list(re.finditer(r"[.!?][\"’”']?(?=\s|$)", complete))
    if sentence_ends:
        complete = complete[:sentence_ends[-1].end()].rstrip()
    else:
        complete = ""
    notes.append(("warn", f"{where} ({who}): the source statement was marked as incomplete. Only its complete sentences will be printed; check the original response."))
    return complete


def load_items(path, profile, limit=None):
    load_fonts()
    notes, items = [], []
    fname = os.path.basename(path)
    headers, rows = read_table(path)
    if not headers:
        return [], [("err", f"{fname}: the spreadsheet is empty.")]
    idx, det = resolve(profile, headers)
    fields, auc = profile["fields"], profile["auction"]

    # Auction status is optional. If the export has no explicit auction column,
    # leave the indicator off rather than guessing that items are or aren't for sale.
    lost = [FIELD_NAME[k] for k in FIELD_KEYS
            if fields.get(k) and idx[k] is None and k != "auction"]
    lost += ["extra detail “%s”" % short(c, 30) for _, i, c in det if c and i is None]
    if idx["artist"] is None and idx["title"] is None:
        return [], [("err", f"{fname}: can't find the artist or title columns - this spreadsheet doesn't match the "
                            f"profile “{profile['profile_name']}”.  Run configure to fix or make a new profile.")]
    for name in lost:
        if name == FIELD_NAME["permission"]:
            explanation = "the generator cannot check which artists have asked for their names to be hidden, so please verify name permissions before printing."
        else:
            explanation = "that information is left off the labels."
        notes.append(("warn", f"{fname}: couldn't find the column for {name} - {explanation}"))

    for rownum, row in rows:
        def g(i):
            return row[i] if i is not None and i < len(row) else None

        source_artist = clean_text(g(idx["artist"]), multi=False)
        artist = source_artist
        title = clean_text(g(idx["title"]), multi=False)
        if not source_artist and not title:
            continue
        where = f"{fname}, row {rownum}"
        who = source_artist or title
        details = []
        for label, i, _ in det:
            v = clean_text(g(i))
            if v and norm(v) not in NA_VALUES:
                details.append((label.strip() + " " + v) if label.strip() else v)
        status = None
        if idx["auction"] is not None:
            ans = clean_text(g(idx["auction"]))
            if ans:
                status = is_yes(ans, auc["yes_words"])
            else:
                notes.append(("warn", f"{where} ({who}): auction status is blank, so no auction marker will be printed."))

        permission_text = clean_text(g(idx["permission"])) if idx["permission"] is not None else ""
        hide_name = _hide_artist_name(permission_text)
        if hide_name:
            if norm(artist) != "anonymous":
                artist = ""
                notes.append(("info", f"{where} ({who}): artist name omitted from the label to respect the display-permission response."))
        elif idx["permission"] is not None and not is_yes(permission_text, ["yes"]):
            notes.append(("warn", f"{where} ({who}): display permission isn't a clear Yes ({short(permission_text or 'blank', 40)}) - please check."))

        statement_text = _prepare_statement_for_label(g(idx["statement"]), where, who, notes)
        item = dict(artist=artist, _artist_key=norm(source_artist), role=clean_text(g(idx["role"])),
                    title=title, details=details, statement=statement_text, auction=status, where=where)
        bad = unsupported(" ".join([artist, item["role"], title, item["statement"]] + details))
        if bad:
            notes.append(("warn", f"{where} ({who}): contains characters this font can't draw: {' '.join(bad[:8])}"))
        items.append(item)
        if limit and len(items) >= limit:
            break
    return items, notes


# ════════════════════════════════════════════════════════════════════════════
#  Drawing the labels
# ════════════════════════════════════════════════════════════════════════════
def make_styles(size):
    def st(name, font, sz, lead, col=(0, 0, 0)):
        return ParagraphStyle(name, fontName=font, fontSize=sz, leading=lead, textColor=colors.Color(*col))
    return dict(name=st("n", FONTS["bold"], 20, 24), role=st("r", FONTS["reg"], 11, 14, (.4, .4, .4)),
                title=st("t", FONTS["ital"], 16, 20), detail=st("d", FONTS["reg"], 12, 15, (.2, .2, .2)),
                stmt=st("s", FONTS["reg"], size, round(size * 1.28, 2), (.15, .15, .15)))


def _quoted(t):
    return t if re.match(r"^[\"“‘']", t) else "“%s”" % t


class Card:
    """One label's content, measured for a given width and statement font size."""

    def __init__(self, item, w, size, ind, stmt=None, show_stmt=True, compact=False):
        self.item, self.w, self.ind, self.blocks = item, w, ind, []
        iw, S = w - 2 * PAD, make_styles(size)

        def add(kind, para, space, width=None):
            _, h = para.wrap(width or iw, 1e7)
            self.blocks.append((kind, space, para, h))

        if item["artist"]:
            add("name", Paragraph(xml_escape(item["artist"]), S["name"]), 0, iw - (ind["w"] if ind else 0))
        if compact:
            add("role", Paragraph("(continued from the previous page)", S["role"]), 2)
        elif item["role"]:
            add("role", Paragraph(xml_escape(item["role"]), S["role"]), 2)
        if item["title"]:
            add("title", Paragraph(xml_escape(_quoted(item["title"])), S["title"]), 8 if self.blocks else 0)
        if not compact:
            for n, d in enumerate(item["details"]):
                add("detail", Paragraph(xml_escape(d), S["detail"]), 6 if n == 0 else 2)
        if show_stmt:
            if stmt is None and item["statement"]:
                stmt = Paragraph(stmt_markup(item["statement"]), S["stmt"])
            if stmt is not None:
                add("stmt", stmt, RULE_GAP)
        self.h = 2 * PAD + sum(sp + h for _, sp, _, h in self.blocks)


def fit_card(item, w, spref, smin, ind, target_h):
    """Shrink the statement text (never below the minimum) until it fits; otherwise the label just grows."""
    size = spref
    while True:
        card = Card(item, w, size, ind)
        if card.h <= target_h or size <= smin:
            return card
        size = max(smin, size - 0.5)


def draw_card(c, card, x, top, h, foot=None):
    w = card.w
    c.setStrokeColorRGB(.55, .55, .55)
    c.setLineWidth(.75)
    c.rect(x, top - h, w, h, stroke=1, fill=0)
    y, name_top = top - PAD, top - PAD
    for kind, sp, p, ph in card.blocks:
        y -= sp
        if kind == "name":
            name_top = y
        if kind == "stmt":
            c.setStrokeColorRGB(.8, .8, .8)
            c.setLineWidth(.5)
            c.line(x + PAD, y + sp / 2, x + w - PAD, y + sp / 2)
        p.drawOn(c, x + PAD, y - ph)
        y -= ph
    if card.ind:                                  # auction dot (green = for auction, red = not)
        r, cx, cy = 7, x + w - PAD - 7, name_top - 12
        col = (.1, .6, .1) if card.ind["status"] else (.8, .1, .1)
        c.setFillColorRGB(*col)
        c.circle(cx, cy, r, stroke=0, fill=1)
        if card.ind["text"]:
            c.setFont(FONTS["bold"], 10)
            c.drawRightString(cx - r - 8, cy - 3.5, card.ind["text"])
    if foot:
        c.setFillColorRGB(.4, .4, .4)
        c.setFont(FONTS["ital"], 9)
        c.drawRightString(x + w - PAD, top - h + 6, foot)


def render_pdf(items, profile, out_path):
    """Arrange labels in cut-friendly rows with aligned borders.

    Labels of similar height share a row, so their top and bottom borders line
    up. If a label is much taller than its neighbours, it gets a full-width
    row instead of stretching a short label beside it. Rows are packed by size
    to use the page efficiently; each row can be cut as a horizontal strip, and
    two-column strips can then be cut down the middle.
    """
    load_fonts()
    lay = {**DEFAULT_LAYOUT, **profile["layout"]}
    pw, ph = PAGE_SIZES.get(lay["page_size"], A4)
    pw, ph = (max(pw, ph), min(pw, ph)) if lay["orientation"] == "landscape" else (min(pw, ph), max(pw, ph))
    m = float(lay["margin_mm"]) * mm
    uw, uh = pw - 2 * m, ph - 2 * m
    cols, rows = max(1, int(lay["columns"])), max(1, int(lay["rows"]))
    cw = uw / cols
    standard_h = uh / rows
    smin = max(MIN_FONT, float(lay["statement_min"]))
    spref = max(smin, float(lay["statement_size"]))
    auc = profile["auction"]
    gap = 8  # keep a small gap between cut rows

    def indicator(it):
        if it["auction"] is None:
            return None
        text = auc["text_yes"] if it["auction"] else auc["text_no"]
        tw = pdfmetrics.stringWidth(text, FONTS["bold"], 10) if text else 0
        return dict(status=bool(it["auction"]), text=text, w=22 + (tw + 8 if text else 0))

    c = canvas.Canvas(out_path, pagesize=(pw, ph))
    c.setTitle(profile["profile_name"] + " - Labels")
    c.setAuthor("Exhibition Label Studio")
    st = dict(started=False, pages=0, grown=0, wide=0, split=0, labels=0)

    def new_page():
        if st["started"]:
            c.showPage()
        st["started"] = True
        st["pages"] += 1

    def split_across_pages(it):
        """Give a statement that cannot fit on one page its own page sequence."""
        st["split"] += 1
        new_page()
        iw, ind = uw - 2 * PAD, indicator(it)
        pending = Paragraph(stmt_markup(it["statement"]), make_styles(smin)["stmt"])
        first = True
        while pending is not None:
            head = Card(it, uw, smin, ind if first else None, show_stmt=False, compact=not first)
            room = max(24, uh - head.h - RULE_GAP)
            pending.wrap(iw, 1e7)
            if pending.height <= room:
                part, rest = pending, None
            else:
                pieces = pending.split(iw, max(12, room - 18))
                if len(pieces) >= 2:
                    part, rest = pieces[0], pieces[1]
                elif pieces:
                    part, rest = pieces[0], None
                else:
                    # An unusually long unbreakable string cannot be split by Paragraph.
                    plain = pending.getPlainText()
                    cut = max(1, min(len(plain), int(len(plain) * 0.6)))
                    part = Paragraph(xml_escape(plain[:cut]), make_styles(smin)["stmt"])
                    rest = Paragraph(xml_escape(plain[cut:]), make_styles(smin)["stmt"]) if plain[cut:] else None
            card = Card(it, uw, smin, ind if first else None, stmt=part, compact=not first)
            card_h = min(uh, card.h)
            draw_card(c, card, m, ph - m, card_h,
                      foot="continued on the next page  ►" if rest else None)
            pending, first = rest, False
            if pending is not None:
                new_page()

    # Only labels that would take up a genuinely large chunk of the page
    # become full-width. Ordinary labels can be taller than the nominal grid
    # row; the height-matching step below pairs them with similarly sized ones.
    if cols == 1:
        tall_row_limit = uh
    else:
        tall_row_limit = min(standard_h * 1.28, uh * 0.58)

    regular = []
    row_objects = []
    split_items = []

    for order, it in enumerate(items):
        ind = indicator(it)
        card = fit_card(it, cw, spref, smin, ind, uh)

        if card.h > uh:
            # Try wrapping a very long label across the page before splitting it.
            wide_card = fit_card(it, uw, spref, smin, ind, uh) if cols > 1 else card
            if wide_card.h <= uh:
                row_objects.append({
                    "cards": [(0, wide_card)], "height": wide_card.h,
                    "wide": cols > 1, "order": order,
                })
            else:
                split_items.append(it)
            continue

        # Keep taller labels in the matching pool too. Two long labels can
        # share a cut row safely when their heights are close; a long label
        # with no suitable partner will be made full-width below.
        regular.append({"card": card, "order": order, "height": card.h,
                        "item": it, "indicator": ind})

    # Pair labels only when their natural heights are close enough that the
    # shorter box will not be stretched by much. Unmatched labels get a half-
    # width row of their own, with no forced extra height.
    regular.sort(key=lambda item: (item["height"], item["order"]))
    i = 0
    while i < len(regular):
        group = regular[i:i + cols]
        can_pair = len(group) == cols
        if can_pair and cols > 1:
            heights = [g["height"] for g in group]
            difference = max(heights) - min(heights)
            tolerance = max(18.0, min(heights) * 0.30)
            can_pair = difference <= tolerance
        if can_pair:
            # Keep the original left-to-right order within each pair/row.
            group.sort(key=lambda item: item["order"])
            row_h = max(g["height"] for g in group)
            row_objects.append({
                "cards": [(col, entry["card"]) for col, entry in enumerate(group)],
                "height": row_h, "wide": False,
                "order": min(entry["order"] for entry in group),
            })
            i += cols
        else:
            entry = regular[i]
            if cols > 1 and entry["height"] > tall_row_limit:
                # A tall label without a similarly sized partner gets its own
                # full-width row, rather than stretching or risking a bad cut.
                wide_card = fit_card(entry["item"], uw, spref, smin,
                                     entry["indicator"], uh)
                if wide_card.h <= uh:
                    row_objects.append({
                        "cards": [(0, wide_card)], "height": wide_card.h,
                        "wide": True, "order": entry["order"],
                    })
                else:
                    split_items.append(entry["item"])
            else:
                row_objects.append({
                    "cards": [(0, entry["card"])], "height": entry["height"],
                    "wide": False, "order": entry["order"],
                })
            i += 1

    # Best-fit decreasing: put the biggest rows first, then place each into the
    # page with the least leftover space that can still fit it. This reduces
    # wasted paper while keeping every row's borders aligned.
    row_objects.sort(key=lambda row: (-row["height"], row["order"]))
    page_rows, page_used = [], []
    for row in row_objects:
        h = row["height"]
        choices = []
        for pi, used in enumerate(page_used):
            extra = (gap if page_rows[pi] else 0) + h
            if used + extra <= uh + 0.01:
                choices.append((uh - (used + extra), pi, extra))
        if choices:
            _, pi, extra = min(choices)
            page_rows[pi].append(row)
            page_used[pi] += extra
        else:
            page_rows.append([row])
            page_used.append(h)

    # Draw each page by full horizontal rows. For paired labels, both boxes use
    # the same row height so the guillotine can cut straight across the sheet.
    for plan in page_rows:
        new_page()
        plan.sort(key=lambda row: (-row["height"], row["order"]))
        y = 0.0
        for row in plan:
            top = ph - m - y
            row_h = row["height"]
            for col, card in row["cards"]:
                x = m if row["wide"] else m + col * cw
                draw_card(c, card, x, top, row_h)
                st["labels"] += 1
                st["grown"] += row_h > card.h + 0.5
            if row["wide"]:
                st["wide"] += 1
            y += row_h + gap

    # Extra-long statements need a dedicated page sequence and are kept away
    # from all other labels so their continuation pages stay easy to cut.
    for it in split_items:
        split_across_pages(it)
        st["labels"] += 1

    if st["started"]:
        c.save()
    return st


def build_labels(profile, files, out_path, limit=None):
    """Read every file, tidy, de-duplicate, sort and draw.  Returns (stats or None, notes)."""
    items, notes = [], []
    for f in files:
        try:
            its, nts = load_items(f, profile, limit)
        except Exception as e:
            notes.append(("err", f"Couldn't read {os.path.basename(f)}: {e}"))
            continue
        items += its
        notes += nts
    seen, unique = set(), []
    for it in items:
        key = (norm(it.get("_artist_key", it["artist"])), norm(it["title"]), norm(it["statement"]))
        if key in seen:
            notes.append(("info", f"{it['where']}: identical entry already included - skipped the duplicate."))
            continue
        seen.add(key)
        unique.append(it)
    sort = profile["layout"].get("sort", "file")
    if sort == "artist":
        unique.sort(key=lambda i: norm(i["artist"]))
    elif sort == "title":
        unique.sort(key=lambda i: norm(i["title"]))
    if limit:
        unique = unique[:limit]
    if not unique:
        return None, notes

    # ── Australian-English spell check (interactive) ──────────────────────
    run_spell_check(unique, notes)

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    try:
        stats = render_pdf(unique, profile, out_path)
    except PermissionError:                                   # PDF is open in a viewer
        root, ext = os.path.splitext(out_path)
        out_path = f"{root} ({datetime.now():%H-%M-%S}){ext}"
        notes.append(("warn", "The old PDF was open in another program, so a new copy was saved beside it."))
        stats = render_pdf(unique, profile, out_path)
    stats["path"] = out_path
    return stats, notes


# ════════════════════════════════════════════════════════════════════════════
#  Terminal look & feel  (plain Python - no extra packages, Windows + Linux)
# ════════════════════════════════════════════════════════════════════════════
IS_TTY = sys.stdin.isatty() and sys.stdout.isatty()
COLOR = sys.stdout.isatty() and "NO_COLOR" not in os.environ
ACC, VIO, DIM = "38;5;80", "38;5;141", "38;5;245"
GOOD, WARN, BAD, BOLD, REV = "38;5;78", "38;5;221", "38;5;203", "1", "7"
SELECT = "48;5;61;38;5;231;1"
_ANSI = re.compile(r"\x1b\[[0-9;?]*[A-Za-z]")


def sty(text, *codes):
    return "\x1b[%sm%s\x1b[0m" % (";".join(codes), text) if COLOR and codes else str(text)


def vlen(s):
    return len(_ANSI.sub("", s))


def term_size():
    s = shutil.get_terminal_size((100, 30))
    return s.columns, s.lines


def W():
    return max(50, min(term_size()[0] - 2, 92))


def wrap_text(text, width):
    out = []
    for para in str(text).split("\n"):
        out += textwrap.wrap(para, max(10, width)) or [""]
    return out


def enable_vt():
    if IS_WIN:
        try:
            import ctypes
            k = ctypes.windll.kernel32
            h, mode = k.GetStdHandle(-11), ctypes.c_uint32()
            if k.GetConsoleMode(h, ctypes.byref(mode)):
                k.SetConsoleMode(h, mode.value | 0x0004)
        except Exception:
            pass


def banner():
    tw, th = term_size()
    if th < 24 or tw < 70:
        return ["  " + sty("░▒▓█ ", VIO) + sty("EXHIBITION LABEL STUDIO", BOLD, ACC) + sty("  gallery labels made easy", DIM)]
    w = min(tw - 2, 76)

    def row(content):
        return "  " + sty("║", VIO) + content + " " * (w - 2 - vlen(content)) + sty("║", VIO)
    title = "  " + sty("░▒▓█", VIO) + "  " + sty("E X H I B I T I O N   L A B E L   S T U D I O", BOLD, ACC) + "  " + sty("█▓▒░", VIO)
    return ["  " + sty("╔" + "═" * (w - 2) + "╗", VIO), row(title),
            row("       " + sty("Beautiful gallery labels, made from your entry form", DIM)),
            "  " + sty("╚" + "═" * (w - 2) + "╝", VIO)]


# ── keyboard ────────────────────────────────────────────────────────────────
_TUI = {"on": False, "old": None}
_SEQ = {b"[A": "up", b"OA": "up", b"[B": "down", b"OB": "down", b"[C": "right", b"OC": "right", b"[D": "left",
        b"OD": "left", b"[H": "home", b"OH": "home", b"[1~": "home", b"[F": "end", b"OF": "end", b"[4~": "end",
        b"[3~": "delete", b"[5~": "pgup", b"[6~": "pgdn"}


def tui_start():
    if _TUI["on"]:
        return
    enable_vt()
    if not IS_WIN:
        _TUI["old"] = termios.tcgetattr(0)
        tty.setcbreak(0)
    sys.stdout.write("\x1b[?1049h\x1b[?25l")
    sys.stdout.flush()
    _TUI["on"] = True


def tui_stop():
    if not _TUI["on"]:
        return
    sys.stdout.write("\x1b[?25h\x1b[?1049l")
    sys.stdout.flush()
    if not IS_WIN and _TUI["old"]:
        termios.tcsetattr(0, termios.TCSADRAIN, _TUI["old"])
    _TUI["on"] = False


atexit.register(tui_stop)


def read_key():
    if IS_WIN:
        ch = msvcrt.getwch()
        if ch in ("\x00", "\xe0"):
            return {"H": "up", "P": "down", "K": "left", "M": "right", "G": "home", "O": "end", "I": "pgup",
                    "Q": "pgdn", "S": "delete"}.get(msvcrt.getwch(), "")
        return {"\r": "enter", "\x1b": "esc", "\x08": "backspace", "\x7f": "backspace", "\t": "tab"}.get(ch, ch)
    fd = sys.stdin.fileno()
    b = os.read(fd, 1)
    if not b:
        raise EOFError
    if b == b"\x1b":
        if not select.select([fd], [], [], 0.05)[0]:
            return "esc"
        seq = os.read(fd, 1)
        if seq in (b"[", b"O"):
            while select.select([fd], [], [], 0.05)[0]:
                ch = os.read(fd, 1)
                seq += ch
                if 0x40 <= ch[0] <= 0x7E:
                    break
        return _SEQ.get(seq, "")
    if b in (b"\r", b"\n"):
        return "enter"
    if b in (b"\x7f", b"\x08"):
        return "backspace"
    if b == b"\t":
        return "tab"
    if b[0] >= 0xC0:
        need = 1 if b[0] < 0xE0 else 2 if b[0] < 0xF0 else 3
        while need:
            b += os.read(fd, 1)
            need -= 1
    return b.decode("utf-8", "replace")


def paint(lines):
    _, th = term_size()
    sys.stdout.write("\x1b[H" + "".join(l + "\x1b[K\n" for l in lines[: th - 1]) + "\x1b[J")
    sys.stdout.flush()


# ── widgets ─────────────────────────────────────────────────────────────────
def menu(title, items, subtitle="", crumb="", default=0, esc="Back"):
    """Arrow-key menu.  items: dicts with label, desc, tag, tag_style, sep.  Returns chosen index or None (Esc)."""
    pick = [k for k, it in enumerate(items) if not it.get("sep")]
    if not pick:
        return None
    i = default if 0 <= default < len(items) and not items[default].get("sep") else pick[0]
    top = 0
    while True:
        tw, th = term_size()
        w = W()
        head = banner() + [""]
        if crumb:
            head.append("  " + sty(crumb, DIM))
        head += ["  " + sty(l, BOLD, ACC) for l in wrap_text(title, w - 4)]
        head += ["  " + l for l in wrap_text(subtitle, w - 4)] if subtitle else []
        head.append("")
        desc = wrap_text(items[i].get("desc", ""), w - 8)[:4] if items[i].get("desc") else []
        foot = [""] + ["   " + sty("│ ", VIO) + sty(l, DIM) for l in desc] + [""]
        foot.append("  " + sty("↑ ↓  choose      Enter  select      Esc  " + esc.lower(), DIM))
        vis = max(3, min(len(items), th - len(head) - len(foot) - 2))
        top = min(max(top, i - vis + 1), i) if i >= 0 else 0
        top = max(0, min(top, len(items) - vis))
        tagw = min(34, max([len(it.get("tag", "")) for it in items] + [0]))
        lw = w - 8 - (tagw + 2 if tagw else 0)
        body = []
        for k in range(top, min(len(items), top + vis)):
            it = items[k]
            if it.get("sep"):
                body.append("  " + sty(("── " + it.get("label", "") + " ").ljust(w - 4, "─"), DIM))
                continue
            label = short(it["label"], lw).ljust(lw)
            tag = short(it.get("tag", ""), tagw).rjust(tagw)
            if k == i:
                body.append(" " + sty(" ►  %s%s " % (label, ("  " + tag) if tagw else ""), SELECT))
            else:
                body.append("     " + label + (("  " + sty(tag, *(it.get("tag_style") or (DIM,)))) if tagw else ""))
        if len(items) > vis:
            foot[-1] += sty("      (%d of %d)" % (pick.index(i) + 1 if i in pick else 0, len(pick)), DIM)
        paint(head + body + foot)
        key = read_key()
        step = {"up": -1, "down": 1}.get(key)
        if step:
            j = i
            while True:
                j = (j + step) % len(items)
                if not items[j].get("sep"):
                    break
            i = j
        elif key == "home":
            i = pick[0]
        elif key == "end":
            i = pick[-1]
        elif key in ("pgup", "pgdn"):
            j = max(0, min(len(items) - 1, i + (vis if key == "pgdn" else -vis)))
            i = min(pick, key=lambda p: abs(p - j))
        elif key == "enter":
            return i
        elif key == "esc":
            return None


def ask_text(title, prompt, default="", help="", crumb="", validate=None):
    """Single-line text box.  Returns the text, or None if Esc."""
    buf, pos, err = list(default), len(default), ""
    while True:
        w = W()
        avail = w - 12
        start = max(0, pos - avail + 1)
        vis = buf[start:start + avail]
        p = pos - start
        text = "".join(vis[:p]) + sty(vis[p] if p < len(vis) else " ", REV) + "".join(vis[p + 1:])
        lines = banner() + [""]
        if crumb:
            lines.append("  " + sty(crumb, DIM))
        lines += ["  " + sty(l, BOLD, ACC) for l in wrap_text(title, w - 4)]
        lines += ["  " + l for l in wrap_text(help, w - 4)] if help else []
        lines += ["", "  " + prompt, "  " + sty("┌" + "─" * (w - 6) + "┐", VIO),
                  "  " + sty("│ ", VIO) + text + " " * (w - 8 - min(len(vis) + 1, avail)) + sty(" │", VIO),
                  "  " + sty("└" + "─" * (w - 6) + "┘", VIO)]
        if err:
            lines += ["", "  " + sty("× " + err, BAD)]
        lines += ["", "  " + sty("Type your answer     Enter  confirm      Esc  cancel", DIM)]
        paint(lines)
        k = read_key()
        if k == "enter":
            val = "".join(buf).strip()
            err = validate(val) if validate else ""
            if not err:
                return val
        elif k == "esc":
            return None
        elif k == "backspace" and pos:
            del buf[pos - 1]
            pos -= 1
        elif k == "delete" and pos < len(buf):
            del buf[pos]
        elif k == "left":
            pos = max(0, pos - 1)
        elif k == "right":
            pos = min(len(buf), pos + 1)
        elif k == "home":
            pos = 0
        elif k == "end":
            pos = len(buf)
        elif len(k) == 1 and k.isprintable():
            buf.insert(pos, k)
            pos += 1


def notice(kind, title, lines, crumb=""):
    col = {"ok": GOOD, "warn": WARN, "err": BAD, "info": ACC}[kind]
    w = W()
    body = []
    for l in lines:
        body += ["   " + x for x in wrap_text(l, w - 8)]
    paint(banner() + [""] + (["  " + sty(crumb, DIM)] if crumb else []) + ["  " + sty(title, BOLD, col), ""] + body +
          ["", "  " + sty("Press Enter to continue", DIM)])
    while read_key() not in ("enter", "esc", " "):
        pass


def confirm(title, question, yes="Yes", no="No", default_yes=True, crumb=""):
    r = menu(title, [dict(label=yes), dict(label=no)], subtitle=question, crumb=crumb, default=0 if default_yes else 1)
    return None if r is None else r == 0


def choose(title, options, current, subtitle="", crumb=""):
    """options = [(label, value, description)] -> value or None"""
    items = [dict(label=l, desc=d, tag="● current" if v == current else "", tag_style=(GOOD,)) for l, v, d in options]
    default = next((k for k, (_, v, _) in enumerate(options) if v == current), 0)
    r = menu(title, items, subtitle=subtitle, crumb=crumb, default=default, esc="Cancel")
    return None if r is None else options[r][1]


def open_path(path):
    try:
        if IS_WIN:
            os.startfile(path)
        else:
            subprocess.Popen(["open" if sys.platform == "darwin" else "xdg-open", path],
                             stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    except Exception:
        pass


# ════════════════════════════════════════════════════════════════════════════
#  Australian English Spell Checker (SymSpell-backed, with built-in fallback)
# ════════════════════════════════════════════════════════════════════════════
_SPELL_WORDS = None  # loaded lazily
_SPELL_HAS_FULL_DICT = False
_SPELL_FULL_DICT_WORD_COUNT = 0
_SPELL_WORDS_BY_LENGTH = None  # used by the built-in fallback suggestion engine
_SPELL_WORDS_INDEX_ID = None
_SPELL_SUGGEST_CACHE = {}
_SPELL_CUSTOM_WORDS = set()

# SymSpell is a small, established spell-suggestion library. It builds an index
# once, so we don't have to compare a typo against every word in a large list.
# If it can't be installed (for example, when offline), the existing checker
# below still works as a fallback.
_SPELL_SYMSPELL = None
_SPELL_SYMSPELL_VERBOSITY = None
_SPELL_SYMSPELL_WORD_SET_ID = None
_SPELL_SYMSPELL_ATTEMPTED = False
_SPELL_SYMSPELL_LOAD_FAILED = False
# A short or broken file should not switch on full-dictionary mode and cause
# every unfamiliar name or art term to be flagged as a spelling mistake.
_SPELL_MIN_FULL_DICT_WORDS = 500

# These are words that often get missed by a general English word list. The
# correction is suggested, not silently applied, so the user stays in control.
_SPELL_COMMON_TYPO = {
    "teh": "the", "taht": "that", "thier": "their", "recieve": "receive",
    "recieved": "received", "beleive": "believe", "wierd": "weird",
    "freind": "friend", "becuase": "because", "definately": "definitely",
    "definetely": "definitely", "seperate": "separate", "occured": "occurred",
    "occurence": "occurrence", "untill": "until", "wich": "which",
    "enviroment": "environment", "begining": "beginning", "comming": "coming",
    "beautifull": "beautiful", "realy": "really", "relevent": "relevant",
    "succesful": "successful", "sucess": "success", "responsability": "responsibility",
    "immediatly": "immediately", "occassion": "occasion", "acheive": "achieve",
    "probaly": "probably", "tomorow": "tomorrow", "grammer": "grammar",
    "arguement": "argument", "goverment": "government", "independant": "independent",
    "knowlege": "knowledge", "maintainance": "maintenance", "neccessary": "necessary",
    "publically": "publicly", "recomend": "recommend", "suprise": "surprise",
    "throught": "through", "tommorow": "tomorrow", "wiches": "wishes",
    "accomodate": "accommodate", "acommodate": "accommodate",
    "adress": "address", "agressive": "aggressive", "apparantly": "apparently",
    "appearence": "appearance", "arguement": "argument", "beleif": "belief",
    "calender": "calendar", "comming": "coming", "commitee": "committee",
    "concious": "conscious", "curiousity": "curiosity", "desparate": "desperate",
    "dissapear": "disappear", "embarass": "embarrass", "existance": "existence",
    "experiance": "experience", "finaly": "finally", "gaurd": "guard",
    "greatful": "grateful", "happend": "happened", "harrass": "harass",
    "heigth": "height", "heirarchy": "hierarchy", "indispensible": "indispensable",
    "lenght": "length", "liason": "liaison", "millenium": "millennium",
    "mispell": "misspell", "neccessarily": "necessarily", "occassionally": "occasionally",
    "priviledge": "privilege", "reccomend": "recommend", "reccommend": "recommend",
    "refered": "referred", "releive": "relieve", "remeber": "remember",
    "restaraunt": "restaurant", "rythm": "rhythm", "succesfully": "successfully",
    "supercede": "supersede", "tommorrow": "tomorrow", "truely": "truly",
    "unfortunatly": "unfortunately", "usefull": "useful", "wether": "whether",
    "wich": "which", "writting": "writing", "yeild": "yield",
}

# Where both forms are seen, prefer the usual Australian spelling for labels.
# Some similar-looking pairs (for example licence/license and practice/practise)
# are deliberately left out because the right spelling depends on grammar.
_SPELL_AU_VARIANTS = {
    "color": "colour", "colors": "colours", "colored": "coloured",
    "coloring": "colouring", "colorful": "colourful", "colorless": "colourless",
    "favorite": "favourite", "favorites": "favourites", "favor": "favour",
    "favors": "favours", "favored": "favoured", "favoring": "favouring",
    "favorable": "favourable", "favorably": "favourably",
    "flavor": "flavour", "flavors": "flavours", "flavored": "flavoured",
    "flavoring": "flavouring", "honor": "honour", "honors": "honours",
    "honored": "honoured", "honoring": "honouring", "honorable": "honourable",
    "neighbor": "neighbour", "neighbors": "neighbours", "neighboring": "neighbouring",
    "neighborhood": "neighbourhood", "labor": "labour", "labors": "labours",
    "labored": "laboured", "laboring": "labouring", "laborer": "labourer",
    "behavior": "behaviour", "behaviors": "behaviours", "behavioral": "behavioural",
    "rumor": "rumour", "rumors": "rumours", "rumored": "rumoured",
    "humor": "humour", "humors": "humours", "humorless": "humourless",
    "vapor": "vapour", "odor": "odour", "odors": "odours", "harbor": "harbour",
    "harbors": "harbours", "tumor": "tumour", "tumors": "tumours",
    "mold": "mould", "molds": "moulds", "molded": "moulded", "molding": "moulding",
    "center": "centre", "centers": "centres", "centered": "centred",
    "centering": "centring", "centerpiece": "centrepiece",
    "theater": "theatre", "theaters": "theatres", "theatergoer": "theatregoer",
    "organize": "organise", "organizes": "organises", "organized": "organised",
    "organizing": "organising", "organization": "organisation",
    "organizations": "organisations", "organizational": "organisational",
    "organizer": "organiser", "organizers": "organisers",
    "recognize": "recognise", "recognizes": "recognises", "recognized": "recognised",
    "recognizing": "recognising", "recognizable": "recognisable",
    "realize": "realise", "realizes": "realises", "realized": "realised",
    "realizing": "realising", "realization": "realisation", "realizations": "realisations",
    "summarize": "summarise", "summarizes": "summarises", "summarized": "summarised",
    "summarizing": "summarising", "memorize": "memorise", "memorized": "memorised",
    "memorizing": "memorising", "visualize": "visualise", "visualized": "visualised",
    "visualizing": "visualising", "customize": "customise", "customized": "customised",
    "customizing": "customising", "modernize": "modernise", "modernized": "modernised",
    "modernizing": "modernising", "utilize": "utilise", "utilized": "utilised",
    "utilizing": "utilising", "standardize": "standardise", "standardized": "standardised",
    "standardizing": "standardising", "centralize": "centralise", "centralized": "centralised",
    "centralizing": "centralising", "categorize": "categorise", "categorized": "categorised",
    "categorizing": "categorising", "capitalize": "capitalise", "capitalized": "capitalised",
    "capitalizing": "capitalising", "civilization": "civilisation", "civilizations": "civilisations",
    "symbolize": "symbolise", "symbolized": "symbolised", "symbolizing": "symbolising",
    "stabilize": "stabilise", "stabilized": "stabilised", "stabilizing": "stabilising",
    "legalize": "legalise", "legalized": "legalised", "legalizing": "legalising",
    "industrialize": "industrialise", "industrialized": "industrialised",
    "industrializing": "industrialising",
    "analyze": "analyse", "analyzes": "analyses", "analyzed": "analysed",
    "analyzing": "analysing", "analyzer": "analyser", "paralyze": "paralyse",
    "paralyzed": "paralysed", "paralyzing": "paralysing",
    "apologize": "apologise", "apologized": "apologised", "apologizing": "apologising",
    "criticize": "criticise", "criticized": "criticised", "criticizing": "criticising",
    "emphasize": "emphasise", "emphasized": "emphasised", "emphasizing": "emphasising",
    "maximize": "maximise", "minimize": "minimise", "prioritize": "prioritise",
    "specialize": "specialise", "specialized": "specialised",
    "defense": "defence", "defenses": "defences", "offense": "offence",
    "offenses": "offences", "gray": "grey", "traveling": "travelling",
    "traveler": "traveller", "travelers": "travellers", "traveled": "travelled",
    "canceling": "cancelling", "canceled": "cancelled", "jewelry": "jewellery",
    "pajamas": "pyjamas", "aluminum": "aluminium", "maneuver": "manoeuvre",
    "maneuvers": "manoeuvres", "mustache": "moustache", "specialty": "speciality",
    "pediatric": "paediatric", "pediatrics": "paediatrics", "pediatrician": "paediatrician",
    "estrogen": "oestrogen", "fetus": "foetus", "anemia": "anaemia",
    "diarrhea": "diarrhoea", "orthopedics": "orthopaedics", "anesthesia": "anaesthesia",
    "cesarean": "caesarean", "esophagus": "oesophagus", "encyclopedia": "encyclopaedia",
    "catalog": "catalogue", "catalogs": "catalogues", "dialog": "dialogue",
    "dialogs": "dialogues", "checkered": "chequered",
    "cozy": "cosy", "cozier": "cosier", "coziest": "cosiest",
    "aging": "ageing", "skeptic": "sceptic",
    "skeptics": "sceptics", "skeptical": "sceptical", "skepticism": "scepticism",
    "esthetic": "aesthetic", "esthetics": "aesthetics", "omelet": "omelette",
    "paleontologist": "palaeontologist", "paleontology": "palaeontology",
    "enroll": "enrol", "enrolls": "enrols",
    "enrollment": "enrolment", "enrollments": "enrolments",
    "fulfill": "fulfil", "fulfills": "fulfils", "fulfillment": "fulfilment",
    "installment": "instalment", "installments": "instalments",
    "skillful": "skilful", "skillfully": "skilfully", "woolen": "woollen",
    "saber": "sabre", "sabers": "sabres", "plow": "plough", "plows": "ploughs",
    "plowed": "ploughed", "plowing": "ploughing",
    "labeling": "labelling", "labeled": "labelled", "modeling": "modelling",
    "modeled": "modelled", "counseling": "counselling", "counseled": "counselled",
    "jeweler": "jeweller", "jewelers": "jewellers", "worshiping": "worshipping",
    "worshiped": "worshipped", "fueling": "fuelling", "fueled": "fuelled",
    "signaling": "signalling", "signaled": "signalled", "traveling": "travelling",
}

# Words that should not be flagged, including common words and known names from
# the example exhibition profile.
_SPELL_KEEP = {
    "a", "i", "ok", "okay", "na", "n/a", "artwork", "artworks",
    "oras", "orana", "oraan", "oranna",
    # Contractions are common in artist statements and should not depend on
    # whether the supplied dictionary happens to contain apostrophe forms.
    "ain't", "aren't", "can't", "couldn't", "could've", "didn't", "doesn't",
    "don't", "hadn't", "hasn't", "haven't", "he'd", "he'll", "he's", "i'd",
    "i'll", "i'm", "i've", "isn't", "it'd", "it'll", "it's", "let's", "mustn't",
    "shan't", "she'd", "she'll", "she's", "shouldn't", "that's", "there's",
    "they'd", "they'll", "they're", "they've", "wasn't", "we'd", "we'll",
    "we're", "we've", "weren't", "what's", "when's", "where's", "who's",
    "won't", "wouldn't", "you'd", "you'll", "you're", "you've",
}

# Proper spellings we always want available, even if the external dictionary
# is missing or a little incomplete.
_BUILTIN_AU_WORDS = set((
    "a an and are as at be been being by can could did do does for from had has have he her hers him his "
    "i if in into is it its may me might mine more most my no not of on one or our ours out over per "
    "she should so some than that the their theirs them then there these they this those through to too "
    "under up us very was we were what when where which who why will with would you your yours "
    "about above across after again against all also although always among another any around art artist "
    "artists artwork beautiful because before best between both bright calm came care change children clay "
    "colour colours colourful centre centres centimetre centimetres favourite favourites neighbour neighbourhood "
    "organise organised organisation recognise recognised realise realised analyse analysed travelled travelling "
    "cancelled cancelling defence offence grey jewellery pyjamas aluminium theatre theatres metre litre mould "
    "paint painting paper people person place plants peaceful picture pictures made make makes making natural "
    "nature new old other outside own part parts people place river sky small soft special story stories "
    "student students texture the time tree trees use used uses using water way well work world year years "
    "environment different family feeling feelings flower flowers garden grass ground heart home house life "
    "light line lines love meaning memory mountain mountains music orange red blue green black white yellow "
    "artist statement visual arts mixed media acrylic pencil pastel sculpture drawing painting"
).split())
_BUILTIN_AU_WORDS.update(_SPELL_AU_VARIANTS.values())
_BUILTIN_AU_WORDS.update(v for v in _SPELL_COMMON_TYPO.values() if " " not in v)

_SPELL_FIELDS = [
    # Names and class labels contain proper nouns and short abbreviations much
    # more often than prose, so checking them creates lots of misleading alerts.
    ("title", "Artwork title"),
    ("statement", "Artist statement"),
]
_SPELL_MIN_LEN = 3
_SPELL_TOKEN_RE = re.compile(r"[^\W\d_]+(?:['’][^\W\d_]+)*", re.UNICODE)


def _normal_spell_word(word):
    return unicodedata.normalize("NFC", str(word)).casefold().replace("’", "'").strip()


_SPELL_DICT_SOURCE = None


def _find_spell_dictionary():
    """Find the main Australian dictionary even if the download kept its suffix."""
    preferred = ["AU_english.txt", "en_AU.txt", "en_AU-ise.txt", "en_AU-large.txt"]
    for filename in preferred:
        path = os.path.join(BASE, filename)
        if os.path.isfile(path):
            return path

    # Browsers sometimes save a download as AU_english (1).txt. Accept that,
    # and prefer the largest matching full dictionary if multiple copies exist.
    candidates = []
    try:
        for filename in os.listdir(BASE):
            lower = filename.casefold()
            if lower.endswith(".txt") and not lower.endswith("_custom.txt") and (
                lower.startswith("au_english") or lower.startswith("en_au")
            ):
                path = os.path.join(BASE, filename)
                if os.path.isfile(path):
                    candidates.append(path)
    except OSError:
        return None
    return max(candidates, key=lambda item: os.path.getsize(item), default=None)


def _load_spell_words():
    """Load the built-in words, a full Australian list and personal additions."""
    global _SPELL_WORDS, _SPELL_HAS_FULL_DICT, _SPELL_FULL_DICT_WORD_COUNT
    global _SPELL_CUSTOM_WORDS, _SPELL_DICT_SOURCE
    if _SPELL_WORDS is not None:
        return _SPELL_WORDS

    words = {_normal_spell_word(w) for w in _BUILTIN_AU_WORDS}
    custom_words = set()
    full_dictionary_words = set()
    main_dictionary = _find_spell_dictionary()
    _SPELL_DICT_SOURCE = os.path.basename(main_dictionary) if main_dictionary else None
    sources = []
    if main_dictionary:
        sources.append((main_dictionary, True))
    custom_path = os.path.join(BASE, "AU_english_custom.txt")
    if os.path.isfile(custom_path):
        sources.append((custom_path, False))

    for path, is_full in sources:
        try:
            with open(path, encoding="utf-8-sig") as fh:
                for line in fh:
                    line = line.strip()
                    if not line or line.startswith("#"):
                        continue
                    # Plain word lists are expected. Strip a possible comment,
                    # trailing punctuation and Hunspell-style /flags if present.
                    word = line.split()[0].split("/", 1)[0].strip(",;:")
                    word = _normal_spell_word(word)
                    if not word or not re.fullmatch(r"[^\W\d_]+(?:['’][^\W\d_]+)*", word, re.UNICODE):
                        continue
                    words.add(word)
                    if is_full:
                        full_dictionary_words.add(word)
                    else:
                        custom_words.add(word)
        except (OSError, UnicodeError):
            # Keep the built-in words available if the file cannot be read.
            continue

    # An empty, truncated or accidentally replaced word list must not make the
    # checker treat all unfamiliar names and art terms as mistakes.
    _SPELL_FULL_DICT_WORD_COUNT = len(full_dictionary_words)
    _SPELL_HAS_FULL_DICT = _SPELL_FULL_DICT_WORD_COUNT >= _SPELL_MIN_FULL_DICT_WORDS
    _SPELL_CUSTOM_WORDS = custom_words
    _SPELL_WORDS = words
    return _SPELL_WORDS

def _levenshtein(a, b):
    """Damerau-Levenshtein distance; a swapped pair counts as one typo."""
    da = {}
    d = [[0] * (len(b) + 2) for _ in range(len(a) + 2)]
    maxdist = len(a) + len(b)
    d[0][0] = maxdist
    for i in range(len(a) + 1):
        d[i + 1][0] = maxdist
        d[i + 1][1] = i
    for j in range(len(b) + 1):
        d[0][j + 1] = maxdist
        d[1][j + 1] = j
    for i in range(1, len(a) + 1):
        db = 0
        for j in range(1, len(b) + 1):
            i2 = da.get(b[j - 1], 0)
            j2 = db
            if a[i - 1] == b[j - 1]:
                d[i + 1][j + 1] = d[i][j]
                db = j
            else:
                d[i + 1][j + 1] = min(d[i][j], d[i + 1][j], d[i][j + 1]) + 1
            d[i + 1][j + 1] = min(
                d[i + 1][j + 1],
                d[i2][j2] + (i - i2 - 1) + 1 + (j - j2 - 1),
            )
        da[a[i - 1]] = i
    return d[len(a) + 1][len(b) + 1]


def _spell_apply_case(suggestion, original):
    """Keep the original word's capitalisation when applying a suggestion."""
    if not suggestion:
        return suggestion
    if original.isupper():
        return suggestion.upper()
    if original.istitle():
        return suggestion[:1].upper() + suggestion[1:]
    if original[:1].isupper():
        return suggestion[:1].upper() + suggestion[1:]
    return suggestion


def _spell_is_likely_proper_name(raw, text, start, lower):
    """Avoid most mid-sentence names while still checking sentence beginnings."""
    if not raw[:1].isupper() or raw.isupper():
        return False
    if lower in _SPELL_COMMON_TYPO or lower in _SPELL_AU_VARIANTS:
        return False
    before = text[:start].rstrip()
    # Treat a capitalised word after a sentence break as normal prose, not a
    # proper name. Also account for quotation marks/brackets before a sentence.
    before = re.sub(r"[\"'“”‘’»«()\[\]{}]+\s*$", "", before).rstrip()
    if not before or before[-1] in ".!?\n":
        return False
    return True


def _spell_known_word(word, word_set):
    """Recognise dictionary words and a few safe, common inflections."""
    w = _normal_spell_word(word)
    if w in word_set or w in _SPELL_KEEP:
        return True
    if w.endswith("'s") and w[:-2] in word_set:
        return True
    if w.endswith("ies") and w[:-3] + "y" in word_set:
        return True
    for suffix in ("s", "es", "ed", "ing", "er", "est", "ly"):
        if not w.endswith(suffix) or len(w) <= len(suffix) + 2:
            continue
        stem = w[:-len(suffix)]
        candidates = {stem, stem + "e"}
        if stem.endswith("i"):
            candidates.add(stem[:-1] + "y")
        # Only allow normal doubled consonants (run -> running, stop -> stopped).
        # In particular, don't let a typo such as "walkking" pass just because
        # removing the extra k happens to reveal a real word.
        if (len(stem) > 2 and stem[-1] == stem[-2]
                and stem[-1] in "bdfglmnprstz" and stem[-3] != stem[-1]):
            candidates.add(stem[:-1])
        if candidates.intersection(word_set):
            return True
    return False


# These common words get a small frequency boost when suggestions are tied.
# SymSpell itself supplies the fast candidate search; this keeps the final order
# sensible even when the Australian source list doesn't include word frequencies.
_SPELL_COMMON_WORD_FREQUENCY = {
    "the": 100, "and": 95, "that": 90, "with": 85, "from": 80,
    "because": 78, "receive": 76, "believe": 74, "beautiful": 72,
    "different": 70, "environment": 68, "necessary": 66, "separate": 64,
    "colour": 62, "centre": 62, "favourite": 60, "organise": 58,
    "recognise": 56, "realise": 54, "travelling": 52, "cancelled": 50,
}


def _load_symspell(word_set):
    """Build the fast suggestion index once, using the user's Australian list."""
    global _SPELL_SYMSPELL, _SPELL_SYMSPELL_VERBOSITY
    global _SPELL_SYMSPELL_WORD_SET_ID, _SPELL_SYMSPELL_ATTEMPTED
    global _SPELL_SYMSPELL_LOAD_FAILED

    if _SPELL_SYMSPELL is not None and _SPELL_SYMSPELL_WORD_SET_ID == id(word_set):
        return _SPELL_SYMSPELL
    if _SPELL_SYMSPELL_ATTEMPTED:
        return None

    _SPELL_SYMSPELL_ATTEMPTED = True
    try:
        from symspellpy import SymSpell, Verbosity
    except ImportError:
        # Match the script's first-run setup style, but keep this dependency
        # optional so the label generator still works without internet access.
        try:
            subprocess.run(
                [sys.executable, "-m", "pip", "install", "--quiet", "symspellpy"],
                check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
                timeout=15,
            )
            importlib.invalidate_caches()
            from symspellpy import SymSpell, Verbosity
        except Exception:
            _SPELL_SYMSPELL_LOAD_FAILED = True
            return None

    try:
        checker = SymSpell(max_dictionary_edit_distance=2, prefix_length=7)
        for candidate in word_set:
            # The current dictionary is a plain list without frequencies. A
            # few everyday words get a boost; all other words start at one.
            if candidate and candidate.isalpha():
                checker.create_dictionary_entry(
                    candidate, _SPELL_COMMON_WORD_FREQUENCY.get(candidate, 1)
                )
    except Exception:
        _SPELL_SYMSPELL_LOAD_FAILED = True
        return None

    _SPELL_SYMSPELL = checker
    _SPELL_SYMSPELL_VERBOSITY = Verbosity
    _SPELL_SYMSPELL_WORD_SET_ID = id(word_set)
    return _SPELL_SYMSPELL


def _spell_suggest(word, word_set, max_dist=2, top_n=8):
    """Find close corrections with SymSpell, then prefer Australian forms."""
    global _SPELL_WORDS_BY_LENGTH, _SPELL_WORDS_INDEX_ID, _SPELL_SYMSPELL_LOAD_FAILED
    w = _normal_spell_word(word)
    cache_key = (id(word_set), w, max_dist, top_n)
    if cache_key in _SPELL_SUGGEST_CACHE:
        return list(_SPELL_SUGGEST_CACHE[cache_key])

    preferred = _SPELL_COMMON_TYPO.get(w) or _SPELL_AU_VARIANTS.get(w)
    if not _SPELL_HAS_FULL_DICT:
        result = [preferred] if preferred else []
        _SPELL_SUGGEST_CACHE[cache_key] = result
        return list(result)

    candidates = {}
    checker = _load_symspell(word_set)
    if checker is not None:
        try:
            matches = checker.lookup(
                w, _SPELL_SYMSPELL_VERBOSITY.ALL,
                max_edit_distance=max_dist, include_unknown=False,
            )
            for match in matches:
                candidate = _normal_spell_word(match.term)
                # Don't suggest American spellings when an Australian option
                # is listed, and don't offer the original word back to itself.
                if candidate in _SPELL_AU_VARIANTS or candidate == w:
                    continue
                if match.distance > 0 and match.distance <= max_dist:
                    candidates[candidate] = match.distance
        except Exception:
            # Unexpected package/API issues shouldn't stop label generation.
            checker = None
            _SPELL_SYMSPELL_LOAD_FAILED = True

    if checker is None:
        # Built-in fallback for offline installs or if SymSpell isn't available.
        if _SPELL_WORDS_BY_LENGTH is None or _SPELL_WORDS_INDEX_ID != id(word_set):
            by_length = {}
            for candidate in word_set:
                if candidate and candidate.isalpha():
                    by_length.setdefault(len(candidate), []).append(candidate)
            _SPELL_WORDS_BY_LENGTH = by_length
            _SPELL_WORDS_INDEX_ID = id(word_set)

        min_len = max(1, len(w) - max_dist)
        max_len = len(w) + max_dist
        for length in range(min_len, max_len + 1):
            for candidate in _SPELL_WORDS_BY_LENGTH.get(length, ()):
                if candidate in _SPELL_AU_VARIANTS or candidate == w:
                    continue
                distance = _levenshtein(candidate, w)
                if 0 < distance <= max_dist:
                    candidates[candidate] = distance

    ordered = sorted(
        candidates,
        key=lambda candidate: (
            candidates[candidate],
            -_SPELL_COMMON_WORD_FREQUENCY.get(candidate, 0),
            candidate,
        ),
    )
    result = []
    if preferred:
        result.append(preferred)
    for candidate in ordered:
        if candidate not in result:
            result.append(candidate)
        if len(result) >= top_n:
            break
    result = result[:top_n]
    _SPELL_SUGGEST_CACHE[cache_key] = result
    return list(result)


def _spell_add_custom_word(word, word_set, notes):
    """Add a one-word spelling to the personal list so future runs remember it."""
    lower = _normal_spell_word(word)
    if not lower or any(ch.isspace() for ch in lower):
        return
    global _SPELL_WORDS_BY_LENGTH, _SPELL_WORDS_INDEX_ID, _SPELL_CUSTOM_WORDS
    word_set.add(lower)
    _SPELL_CUSTOM_WORDS.add(lower)
    _SPELL_WORDS_BY_LENGTH = None
    _SPELL_WORDS_INDEX_ID = None
    _SPELL_SUGGEST_CACHE.clear()
    if _SPELL_SYMSPELL is not None and lower.isalpha():
        try:
            _SPELL_SYMSPELL.create_dictionary_entry(
                lower, _SPELL_COMMON_WORD_FREQUENCY.get(lower, 1)
            )
        except Exception:
            pass
    path = os.path.join(BASE, "AU_english_custom.txt")
    try:
        existing = set()
        if os.path.isfile(path):
            with open(path, encoding="utf-8-sig") as fh:
                existing = {_normal_spell_word(line.strip()) for line in fh if line.strip() and not line.lstrip().startswith("#")}
        if lower not in existing:
            with open(path, "a", encoding="utf-8") as fh:
                fh.write(lower + "\n")
        notes.append(("ok", "Added “%s” to your custom spelling dictionary." % word))
    except (OSError, UnicodeError):
        notes.append(("warn", "Kept “%s” as a recognised word for this run, but couldn't save it to AU_english_custom.txt." % word))


def spell_check_field(field_label, text, word_set, notes, skip_all_ref):
    """Check text and let the user correct, skip or remember questionable words."""
    if not text or skip_all_ref[0]:
        return text, False

    replacements = []
    full_text = str(text)
    for match in _SPELL_TOKEN_RE.finditer(full_text):
        raw = match.group()
        stripped = raw.strip("'’-")
        if len(stripped) < _SPELL_MIN_LEN or any(ch.isdigit() for ch in stripped):
            continue
        if stripped.isupper() and len(stripped) <= 8:
            continue
        lower = _normal_spell_word(stripped)
        if lower in _SPELL_KEEP:
            continue

        # Explicit US-to-Australian preferences and common typo fixes are
        # checked even if the supplied dictionary happens to contain the US form.
        forced_correction = (lower in _SPELL_AU_VARIANTS or lower in _SPELL_COMMON_TYPO) and lower not in _SPELL_CUSTOM_WORDS
        if not forced_correction and _spell_is_likely_proper_name(raw, full_text, match.start(), lower):
            continue
        if not forced_correction and _spell_known_word(lower, word_set):
            continue
        # Without AU_english.txt, only flag high-confidence known problems. This
        # avoids treating every unfamiliar name or art term as a spelling error.
        if not forced_correction and not _SPELL_HAS_FULL_DICT:
            continue
        replacements.append((match.start(), match.end(), raw, stripped, lower))

    if not replacements:
        return text, False

    if not IS_TTY:
        for _, _, raw, stripped, lower in replacements:
            suggestions = _spell_suggest(stripped, word_set)
            if suggestions:
                notes.append(("warn", "Spell check (%s): “%s” may need checking. Suggestion: “%s”." %
                              (field_label, raw, _spell_apply_case(suggestions[0], raw))))
            else:
                notes.append(("warn", "Spell check (%s): “%s” may be misspelled." % (field_label, raw)))
        return text, False

    result_chars = list(full_text)
    offset = 0
    changed = False
    for start, end, raw, stripped, lower in replacements:
        if skip_all_ref[0]:
            break
        adj_start, adj_end = start + offset, end + offset
        current = "".join(result_chars[adj_start:adj_end])
        entire = "".join(result_chars)
        ctx_before = entire[max(0, adj_start - 40):adj_start].strip()
        ctx_after = entire[adj_end:adj_end + 40].strip()
        context_str = ("…" if adj_start > 40 else "") + ctx_before + \
                      "  " + sty(current, BOLD, BAD) + "  " + ctx_after + \
                      ("…" if adj_end + 40 < len(entire) else "")

        suggestions = _spell_suggest(stripped, word_set)
        correct = _spell_correct_one(field_label, current, stripped, context_str,
                                     suggestions, skip_all_ref, word_set, notes)
        if correct is None:
            continue
        if correct == "___ADDED_TO_DICT___":
            continue
        if correct == "___SKIP_ALL___" or skip_all_ref[0]:
            skip_all_ref[0] = True
            break
        if correct != current:
            result_chars[adj_start:adj_end] = list(correct)
            offset += len(correct) - len(current)
            changed = True
            if correct != raw:
                notes.append(("ok", "Spell check (%s): “%s” → “%s”" % (field_label, current, correct)))
    return "".join(result_chars), changed


def _spell_correct_one(field_label, raw, stripped, context_str, suggestions,
                       skip_all_ref, word_set, notes):
    """Show suggestions and actions for one word. Return the chosen replacement."""
    items = []
    if suggestions:
        for suggestion in suggestions[:8]:
            display = _spell_apply_case(suggestion, raw)
            items.append(dict(label=display, value=display, desc=""))
        items.append(dict(sep=True, label="More actions"))
    else:
        items.append(dict(sep=True, label="No suggestions found"))
    items.append(dict(label="[Skip this time]", value="__skip__", desc="Leave “%s” as-is." % raw))
    items.append(dict(label="[Add to dictionary]", value="__add__",
                      desc="Remember this word for future checks."))
    items.append(dict(label="[Skip all remaining]", value="__skip_all__",
                      desc="Don't ask again during this run."))
    items.append(dict(label="[Type a replacement…]", value="__type__",
                      desc="Manually type your own correction."))

    while True:
        header_parts = [sty("Spell Check", BOLD, ACC),
                        sty("Field: ", VIO) + sty(field_label, BOLD), context_str]
        if suggestions:
            header_parts.append(sty("Suggestions (Australian English):", VIO))
        else:
            header_parts.append(sty("No suggestions found for “%s”" % raw, WARN))
        r = menu("Spell Check", items, subtitle="\n".join(header_parts),
                 crumb="Spell Check", esc="Skip this word")
        if r is None:
            return None

        value = items[r]["value"]
        if value == "__skip__":
            return raw
        if value == "__add__":
            _spell_add_custom_word(stripped, word_set, notes)
            return "___ADDED_TO_DICT___"
        if value == "__skip_all__":
            skip_all_ref[0] = True
            return "___SKIP_ALL___"
        if value == "__type__":
            typed = ask_text("Type your correction", "Replacement for “%s”:" % raw,
                             default=raw, crumb="Spell Check",
                             help="Type the correct spelling and press Enter.")
            if typed is not None:
                return typed
            continue
        if not value.startswith("__"):
            return value


def run_spell_check(items, notes):
    """Run the Australian English spell check on label text, modifying items in-place."""
    word_set = _load_spell_words()
    if _SPELL_DICT_SOURCE and _SPELL_DICT_SOURCE != "AU_english.txt":
        notes.append(("info", "Spell checker is using %s because AU_english.txt wasn't found. "
                              "Rename the preferred dictionary to AU_english.txt if you want to make it the default." % _SPELL_DICT_SOURCE))
    if not _SPELL_HAS_FULL_DICT:
        if _SPELL_FULL_DICT_WORD_COUNT:
            reason = "AU_english.txt only has %d usable words" % _SPELL_FULL_DICT_WORD_COUNT
        else:
            reason = "AU_english.txt was not found or could not be read"
        notes.append(("warn", reason + " — checking common typos and known US spellings only. "
                              "Add a full Australian English word list (at least %d words) beside this script "
                              "for broader spell checking." % _SPELL_MIN_FULL_DICT_WORDS))

    tui_was_off = not _TUI["on"]
    if tui_was_off and IS_TTY:
        tui_start()

    total_checked = 0
    total_changed = 0
    skip_all_ref = [False]
    try:
        for item in items:
            if skip_all_ref[0]:
                break
            for key, label in _SPELL_FIELDS:
                if skip_all_ref[0]:
                    break
                text = item.get(key, "")
                if not text:
                    continue
                corrected, changed = spell_check_field(label, text, word_set, notes, skip_all_ref)
                if changed:
                    item[key] = corrected
                    total_changed += 1
                total_checked += 1

            # Extra detail values often include the names of art materials.
            for detail_index, detail_text in enumerate(item.get("details", [])):
                if skip_all_ref[0]:
                    break
                if not detail_text:
                    continue
                corrected, changed = spell_check_field("Detail line %d" % (detail_index + 1),
                                                       detail_text, word_set, notes, skip_all_ref)
                if changed:
                    item["details"][detail_index] = corrected
                    total_changed += 1
                total_checked += 1
    finally:
        if tui_was_off and _TUI["on"]:
            tui_stop()

    if _SPELL_SYMSPELL_LOAD_FAILED:
        notes.append(("info", "Fast spell suggestions weren't available, so the built-in fallback was used. "
                              "The checker can try installing symspellpy again on the next run."))

    if total_changed:
        notes.append(("ok", "Spell checker: corrected spelling in %d field%s out of %d checked field%s." %
                       (total_changed, "" if total_changed == 1 else "s",
                        total_checked, "" if total_checked == 1 else "s")))


# ════════════════════════════════════════════════════════════════════════════
#  Output for "make the labels" (ordinary scrolling text, still pretty)
# ════════════════════════════════════════════════════════════════════════════
def say(kind, text, indent=2):
    mark = {"ok": sty("●", GOOD), "info": sty("●", ACC), "warn": sty("▲", WARN), "err": sty("×", BAD),
            "step": sty("►", VIO)}[kind]
    print(" " * indent + mark + " " + text)


def box(lines, color):
    w = max(vlen(l) for l in lines) + 4
    print("  " + sty("┌" + "─" * w + "┐", color))
    for l in lines:
        print("  " + sty("│", color) + "  " + l + " " * (w - 2 - vlen(l)) + sty("│", color))
    print("  " + sty("└" + "─" * w + "┘", color))


def run_generation(profile, files, limit=None, out_path=None):
    say("step", "Using profile: " + sty(profile["profile_name"], BOLD))
    lay = profile["layout"]
    say("info", "Layout: %s × %s labels per page · %s %s · statement text %g–%g pt"
        % (lay["columns"], lay["rows"], lay["page_size"], lay["orientation"],
           max(MIN_FONT, float(lay["statement_min"])), max(float(lay["statement_min"]), float(lay["statement_size"]))))
    out = out_path or os.path.join(OUTPUT_FOLDER, safe_filename(profile["profile_name"]) + " - Labels.pdf")
    say("step", "Building labels from %d file%s..." % (len(files), "" if len(files) == 1 else "s"))
    stats, notes = build_labels(profile, files, out, limit)
    print()
    errs = [t for k, t in notes if k == "err"]
    warns = [t for k, t in notes if k == "warn"]
    infos = [t for k, t in notes if k == "info"]
    for t in errs:
        say("err", t)
    for t in warns[:15]:
        say("warn", t)
    if len(warns) > 15:
        say("warn", "...and %d more warnings." % (len(warns) - 15))
    for t in infos[:5]:
        say("info", t)
    if errs or warns or infos:
        print()
    if not stats:
        box([sty("No labels were made.", BAD),
             "Check the messages above, or run configure to fix the profile."], BAD)
        return None
    extra = []
    if stats["grown"]:
        extra.append("%d label%s grew taller to fit a long statement" % (stats["grown"], "" if stats["grown"] == 1 else "s"))
    if stats["wide"]:
        extra.append("%d very long one%s got a full-width label" % (stats["wide"], "" if stats["wide"] == 1 else "s"))
    if stats["split"]:
        extra.append("%d extremely long one%s continue%s on a second page" % (stats["split"], "" if stats["split"] == 1 else "s", "s" if stats["split"] == 1 else ""))
    lines = [sty("Done!  %d labels on %d page%s" % (stats["labels"], stats["pages"], "" if stats["pages"] == 1 else "s"), BOLD, GOOD)]
    lines += [sty(e, DIM) for e in extra]
    lines += ["", "Saved to:", sty(os.path.relpath(stats["path"], BASE), ACC)]
    box(lines, GOOD)
    return stats


def pick_profile_for_generate(profiles, heads, wanted):
    if wanted:
        for p in profiles:
            if norm(p["profile_name"]) == norm(wanted):
                return p
        say("err", "No profile called “%s”.  Available: %s" % (wanted, ", ".join(p["profile_name"] for p in profiles)))
        return None
    scores = {p["profile_name"]: (sum(match_score(p, h) for h in heads.values()) / len(heads)) if heads else 0.0
              for p in profiles}
    best = max(range(len(profiles)), key=lambda k: scores[profiles[k]["profile_name"]])
    if len(profiles) == 1 or not IS_TTY:
        return profiles[best]
    tui_start()
    try:
        items = [dict(label=p["profile_name"], tag="%d%% match" % round(scores[p["profile_name"]] * 100),
                      tag_style=(GOOD if scores[p["profile_name"]] > .8 else WARN,),
                      desc="How well this profile's columns fit your spreadsheet(s).") for p in profiles]
        r = menu("Which form profile should be used?", items, default=best, esc="Quit",
                 subtitle="You have more than one profile.  The best match is highlighted.")
    finally:
        tui_stop()
    return None if r is None else profiles[r]


def cmd_generate(args):
    enable_vt()
    print()
    print("\n".join(banner()))
    print()
    os.makedirs(INPUT_FOLDER, exist_ok=True)
    os.makedirs(OUTPUT_FOLDER, exist_ok=True)
    profiles = load_profiles()
    for bad in BAD_PROFILES:
        say("warn", "Skipped unreadable profile file: " + bad)
    if not profiles:
        say("err", "There are no form profiles yet.  Run the configure script to create one.")
        return 1
    files = find_data_files()
    if not files:
        say("warn", "No spreadsheets found.  Drop your .xlsx or .csv files into the folder:")
        say("info", sty("EXCEL_FILES_HERE", ACC, BOLD) + "   and run this again.")
        return 1
    say("ok", "Found %d spreadsheet%s in EXCEL_FILES_HERE" % (len(files), "" if len(files) == 1 else "s"))
    heads = {}
    for f in files:
        try:
            heads[f] = read_table(f)[0]
        except Exception:
            pass
    prof = pick_profile_for_generate(profiles, heads, args.profile)
    if prof is None:
        return 1
    return 0 if run_generation(prof, files, args.limit) else 1


# ════════════════════════════════════════════════════════════════════════════
#  Configure  (the friendly set-up screens)
# ════════════════════════════════════════════════════════════════════════════
SESSION = {"headers": [], "rows": []}


def read_example(path):
    try:
        headers, rows = read_table(path)
    except Exception as e:
        notice("err", "That file couldn't be opened", [str(e), "Make sure it is a .xlsx or .csv file and is not open in another program."])
        return None
    if not headers:
        notice("err", "That spreadsheet looks empty", ["Export the form responses to Excel again and try once more."])
        return None
    SESSION.update(headers=headers, rows=rows)
    return headers, rows


def pick_data_file(title, crumb=""):
    files = find_data_files()
    items = [dict(label=os.path.basename(f), desc="Found in the EXCEL_FILES_HERE folder.", value=f) for f in files]
    items.append(dict(label="Type or paste a file location…", value="__path__",
                      desc="Tip: on most computers you can drag a file into this window to paste its location."))
    r = menu(title, items, crumb=crumb, esc="Cancel",
             subtitle="Choose a spreadsheet exported from your form.  It is only used to read the column names."
             if files else "No spreadsheets are in the EXCEL_FILES_HERE folder yet, so you can paste a location instead.")
    if r is None:
        return None
    if items[r]["value"] != "__path__":
        return items[r]["value"]
    while True:
        s = ask_text("Where is the spreadsheet?", "File location:", crumb=crumb,
                     help="Example:  C:\\Users\\me\\Downloads\\responses.xlsx   or   /home/me/responses.xlsx")
        if s is None:
            return None
        s = os.path.expanduser(s.strip().strip("\"'").replace("\\ ", " "))
        if os.path.isfile(s) and s.lower().endswith(DATA_EXTS):
            return s
        notice("err", "I couldn't use that", ["The file wasn't found, or it isn't an .xlsx / .csv file."])


def sample_value(headers, i):
    if SESSION["headers"] != headers:
        return ""
    for _, row in SESSION["rows"][:10]:
        v = clean_text(row[i]) if i < len(row) else ""
        if v:
            return v
    return ""


def pick_column(prof, key, headers, crumb, optional, suggestion=None, question=""):
    """Column chooser.  Returns column index, -1 for 'leave out', or None for Esc."""
    taken = {k: find_col(headers, v) for k, v in prof["fields"].items() if v and k != key}
    used = {i: FIELD_NAME[k] for k, i in taken.items() if i is not None}
    cur = find_col(headers, prof["fields"].get(key, "")) if key in prof["fields"] else None
    items = []
    if optional:
        items.append(dict(label="— Leave this out —", value=-1, tag="skip", desc="This form doesn't ask for it, or you don't want it on the labels."))
    for i, h in enumerate(headers):
        tag, style = "", (DIM,)
        if i == cur:
            tag, style = "● current", (GOOD,)
        elif i == suggestion:
            tag, style = "★ suggested", (WARN,)
        elif i in used:
            tag = "in use: " + used[i]
        ex = sample_value(headers, i)
        full = re.sub(r"\s+", " ", h.replace("\xa0", " ")).strip()
        desc = ("Example answer:  “%s”" % short(ex, 150) if ex else "") + ("\n" if ex else "") + short(full, 200)
        items.append(dict(label=nice_header(h, 70) or "(blank)", tag=tag, tag_style=style, value=i, desc=desc))
    default = next((k for k, it in enumerate(items) if it["value"] == (cur if cur is not None else suggestion)), 0)
    r = menu(question, items, crumb=crumb, default=default, esc="Cancel",
             subtitle="Pick the column from your spreadsheet.  ★ marks my best guess.")
    return None if r is None else items[r]["value"]


def set_field(prof, key, headers, crumb, optional, suggest=True):
    """Ask for one field.  Returns False if the user pressed Esc."""
    q = next(x[2] for x in FIELDS if x[0] == key)
    guess = guess_column(headers, next(x[4] for x in FIELDS if x[0] == key)) if suggest else None
    r = pick_column(prof, key, headers, crumb, optional, guess, q)
    if r is None:
        return False
    prof["fields"][key] = "" if r == -1 else header_phrase(headers[r])
    return True


def edit_details(prof, headers, crumb):
    while True:
        items = []
        for n, d in enumerate(prof["details"]):
            j = find_col(headers, d["column"])
            items.append(dict(label="Line %d:  %s‹%s›" % (n + 1, (d["label"] + " ") if d.get("label") else "", nice_header(headers[j], 40) if j is not None else d["column"]),
                              tag="" if j is not None else "not found", tag_style=(WARN,), value=("edit", n),
                              desc="Press Enter to change or remove this line."))
        items.append(dict(label="+ Add a line", value=("add", 0), desc="For example the MEDIUM, the SIZE, the PRICE… anything you'd like printed under the title."))
        items.append(dict(label="Done", value=("done", 0)))
        r = menu("Extra detail lines", items, crumb=crumb, default=len(items) - 1, esc="Done",
                 subtitle="These lines are printed under the artwork title, in order.")
        if r is None or items[r]["value"][0] == "done":
            return
        kind, n = items[r]["value"]
        if kind == "add":
            tmp = new_profile("tmp")
            c = pick_column(tmp, "_", headers, crumb, False, None, "Which column should this line show?")
            if c is None:
                continue
            lab = ask_text("Text before the answer", "Label (optional):", crumb=crumb,
                           help="For example  Medium:  or  Size:  - or leave it empty to print just the answer.")
            if lab is None:
                continue
            prof["details"].append({"label": lab, "column": header_phrase(headers[c])})
        else:
            d = prof["details"][n]
            a = choose("What would you like to do?", [("Change the column", "col", ""), ("Change the label text", "lab", ""),
                                                        ("Move up", "up", ""), ("Remove this line", "del", "")], None, crumb=crumb)
            if a == "col":
                tmp = new_profile("tmp")
                c = pick_column(tmp, "_", headers, crumb, False, None, "Which column should this line show?")
                if c is not None:
                    d["column"] = header_phrase(headers[c])
            elif a == "lab":
                lab = ask_text("Text before the answer", "Label (optional):", d.get("label", ""), crumb=crumb)
                if lab is not None:
                    d["label"] = lab
            elif a == "up" and n > 0:
                prof["details"][n - 1], prof["details"][n] = prof["details"][n], prof["details"][n - 1]
            elif a == "del":
                del prof["details"][n]


def layout_menu(prof, crumb):
    L = prof["layout"]
    while True:
        smin = max(MIN_FONT, float(L["statement_min"]))
        items = [
            dict(label="Paper size", tag=L["page_size"], value="page_size"),
            dict(label="Paper direction", tag=L["orientation"].capitalize(), value="orientation"),
            dict(label="Labels across the page", tag=str(L["columns"]), value="columns",
                 desc="More columns = narrower labels."),
            dict(label="Labels down the page", tag=str(L["rows"]), value="rows",
                 desc="This is the standard label height.  Labels with long statements grow taller than this."),
            dict(label="Statement text size", tag="%g pt" % L["statement_size"], value="statement_size",
                 desc="Preferred size.  If a statement doesn't fit, the text shrinks - but never below the smallest size below."),
            dict(label="Smallest statement text allowed", tag="%g pt" % smin, value="statement_min",
                 desc="Text never goes below 10 pt.  Past this, the label grows taller instead."),
            dict(label="Margin around the page", tag="%g mm" % L["margin_mm"], value="margin_mm",
                 desc="Printers can't print right to the edge of the paper, so keep at least 5 mm."),
            dict(label="Order of the labels", tag={"file": "as in spreadsheet", "artist": "artist A–Z", "title": "title A–Z"}[L["sort"]], value="sort"),
            dict(label="Auction wording", tag="“%s”" % prof["auction"]["text_yes"], value="auction",
                 desc="The words printed beside the green dot.  (The red dot can carry words too.)"),
            dict(label="Back", value="back"),
        ]
        r = menu("Label layout & style", items, crumb=crumb, esc="Back")
        if r is None or items[r]["value"] == "back":
            return
        k = items[r]["value"]
        opts = {
            "page_size": [(n, n, "") for n in PAGE_SIZES],
            "orientation": [("Landscape (wide)", "landscape", ""), ("Portrait (tall)", "portrait", "")],
            "columns": [(str(n), n, "") for n in (1, 2, 3)],
            "rows": [(str(n), n, "") for n in (1, 2, 3, 4)],
            "statement_size": [("%d pt" % n, n, "") for n in range(10, 19)],
            "statement_min": [("%d pt" % n, n, "") for n in (10, 11, 12, 14)],
            "margin_mm": [("%d mm" % n, n, "") for n in (5, 8, 10, 12, 15, 20)],
            "sort": [("As in the spreadsheet", "file", ""), ("Artist name A–Z", "artist", ""), ("Artwork title A–Z", "title", "")],
        }
        if k == "auction":
            t = ask_text("Auction wording", "Words beside the GREEN dot:", prof["auction"]["text_yes"], crumb=crumb,
                         help="Leave empty to show just the dot.")
            if t is not None:
                prof["auction"]["text_yes"] = t
            t = ask_text("Auction wording", "Words beside the RED dot:", prof["auction"]["text_no"], crumb=crumb,
                        help="Leave empty to use the default:  Not available for Auction")
            if t is not None:
                prof["auction"]["text_no"] = t
        else:
            v = choose(items[r]["label"], opts[k], L[k], crumb=crumb)
            if v is not None:
                L[k] = v


def make_test_pdf(prof, crumb):
    path = pick_data_file("Make a test PDF - pick a spreadsheet", crumb)
    if not path:
        return
    out = os.path.join(OUTPUT_FOLDER, "PREVIEW - " + safe_filename(prof["profile_name"]) + ".pdf")
    stats, notes = build_labels(prof, [path], out, limit=6)
    if not stats:
        notice("err", "Nothing could be made", [t for _, t in notes] or ["No usable rows were found."], crumb)
        return
    lines = ["A test PDF with the first %d labels was saved to:" % stats["labels"], os.path.relpath(stats["path"], BASE), ""]
    lines += ["▲ " + t for k, t in notes if k in ("warn", "err")][:6]
    lines += ["", "It should open automatically.  If it doesn't, open the Generated_Labels folder."]
    open_path(stats["path"])
    notice("ok", "Test PDF ready", lines, crumb)


def ensure_headers(prof, crumb):
    """The saved form columns - or ask for an example spreadsheet if none are saved yet."""
    if prof["form_columns"]:
        return prof["form_columns"]
    notice("info", "One quick thing first", ["This profile doesn't have a list of the form's columns saved yet.",
                                             "Pick an example spreadsheet from the form so I can show you the columns."], crumb)
    path = pick_data_file("Pick an example spreadsheet", crumb)
    got = read_example(path) if path else None
    if not got:
        return None
    prof["form_columns"] = got[0]
    save_profile(prof)
    return got[0]


def profile_hub(prof):
    while True:
        heads = prof["form_columns"]
        crumb = "Profiles  ›  " + prof["profile_name"]
        items = []
        for k, name, q, req, _ in FIELDS:
            phrase = prof["fields"].get(k, "")
            if not phrase:
                tag, style = "(not used)", (DIM,)
            elif heads and find_col(heads, phrase) is None:
                tag, style = "▲ not in form", (WARN,)
            else:
                j = find_col(heads, phrase) if heads else None
                tag, style = (nice_header(heads[j], 34) if j is not None else short(phrase, 34)), (GOOD,)
            items.append(dict(label=name + (" *" if req else ""), tag=tag, tag_style=style, value=("field", k),
                              desc=q + "   Press Enter to change it."))
        items.append(dict(label="Extra detail lines", tag="%d line%s" % (len(prof["details"]), "" if len(prof["details"]) == 1 else "s"),
                          value=("details", 0), desc="Medium, size, price… printed under the title."))
        L = prof["layout"]
        items += [
            dict(sep=True, label="Look"),
            dict(label="Label layout & style", tag="%s %s · %s×%s" % (L["page_size"], L["orientation"], L["columns"], L["rows"]), value=("layout", 0),
                 desc="Paper size, labels per page, text size, ordering, auction wording."),
            dict(label="Make a test PDF", value=("test", 0), desc="Builds the first few labels from a spreadsheet so you can check how they look."),
            dict(sep=True, label="Manage"),
            dict(label="Read columns from a new spreadsheet", value=("rescan", 0), desc="Use this if the form gained new questions."),
            dict(label="Rename this profile", value=("rename", 0)),
            dict(label="Make a copy of this profile", value=("copy", 0), desc="Handy for next year's exhibition."),
            dict(label="Delete this profile", value=("delete", 0)),
            dict(label="Back to the main menu", value=("back", 0)),
        ]
        r = menu("Profile: " + prof["profile_name"], items, crumb="Profiles", esc="Back",
                 subtitle="Everything is saved automatically.  * = needed on every label.")
        if r is None or items[r]["value"][0] == "back":
            return
        kind, k = items[r]["value"]
        if kind == "field":
            h = ensure_headers(prof, crumb)
            if h:
                req = next(x[3] for x in FIELDS if x[0] == k)
                set_field(prof, k, h, crumb, optional=not req, suggest=False)
        elif kind == "details":
            h = ensure_headers(prof, crumb)
            if h:
                edit_details(prof, h, crumb)
        elif kind == "layout":
            layout_menu(prof, crumb)
        elif kind == "test":
            make_test_pdf(prof, crumb)
        elif kind == "rescan":
            path = pick_data_file("Read columns from a spreadsheet", crumb)
            got = read_example(path) if path else None
            if got:
                prof["form_columns"] = got[0]
        elif kind == "rename":
            n = ask_text("Rename profile", "New name:", prof["profile_name"], crumb=crumb, validate=lambda v: name_problem(v, prof))
            if n:
                prof["profile_name"] = n
        elif kind == "copy":
            n = ask_text("Copy profile", "Name for the copy:", prof["profile_name"] + " (copy)", crumb=crumb, validate=lambda v: name_problem(v, None))
            if n:
                dup = copy.deepcopy(prof)
                dup["profile_name"] = n
                dup.pop("_file", None)
                save_profile(dup)
                notice("ok", "Copy created", ["“%s” is ready.  Choose it from the profile list to edit it." % n], crumb)
        elif kind == "delete":
            if confirm("Delete this profile?", "“%s” will be removed for good.  (Your spreadsheets are not touched.)" % prof["profile_name"],
                       yes="Yes, delete it", no="No, keep it", default_yes=False, crumb=crumb):
                if prof.get("_file") and os.path.exists(prof["_file"]):
                    os.remove(prof["_file"])
                return
        save_profile(prof)


def name_problem(v, current):
    if not v:
        return "Please type a name."
    if not safe_filename(v).strip():
        return "That name can't be used."
    p = profile_path(v)
    if os.path.exists(p) and not (current and current.get("_file") and os.path.abspath(current["_file"]) == os.path.abspath(p)):
        return "A profile with that name already exists."
    return ""


def wizard_new():
    crumb = "New profile"
    notice("info", "Let's teach the app about your form", [
        "A profile remembers what each column of your form means - which one is the artist's name, which is the title, and so on.",
        "You only do this once per form.  I'll guess the answers for you, so most steps are just a matter of pressing Enter.",
        "You'll need an example spreadsheet exported from the form (Forms → Open in Excel).",
    ], crumb)
    path = pick_data_file("Step 1 - Pick an example spreadsheet", crumb + "  ›  Step 1 of 4")
    if not path:
        return
    got = read_example(path)
    if not got:
        return
    headers, rows = got
    name = ask_text("Step 2 - Name this form profile", "Profile name:", "", crumb=crumb + "  ›  Step 2 of 4",
                    help="Something like:  What Orana Means To Me - 2026   (exhibition name, a dash, then the year).",
                    validate=lambda v: name_problem(v, None))
    if not name:
        return
    prof = new_profile(name)
    prof["form_columns"] = headers
    order = [k for k, *_ in FIELDS]
    for n, key in enumerate(order, 1):
        req = next(x[3] for x in FIELDS if x[0] == key)
        if not set_field(prof, key, headers, "%s  ›  Step 3 of 4  ›  Question %d of %d" % (crumb, n, len(order)), optional=not req):
            if confirm("Stop setting up?", "Your answers so far will be thrown away.", yes="Yes, stop", no="No, carry on",
                       default_yes=False, crumb=crumb) in (True, None):
                return
            set_field(prof, key, headers, crumb, optional=not req)
    used = {find_col(headers, v) for v in prof["fields"].values() if v}
    for hints in (["medium", "material"], ["dimension", "size"]):
        g = guess_column(headers, hints, taken=used)
        if g is not None:
            used.add(g)
            prof["details"].append({"label": "", "column": header_phrase(headers[g])})
    edit_details(prof, headers, crumb + "  ›  Step 4 of 4")
    save_profile(prof)
    if confirm("Profile saved!", "“%s” is ready to use.  Would you like to make a test PDF now to see how it looks?" % name,
               yes="Yes, make a test PDF", no="Not now", crumb=crumb):
        make_test_pdf(prof, crumb)
    profile_hub(prof)


def folders_menu():
    r = choose("Open a folder", [("Where I drop my spreadsheets  (EXCEL_FILES_HERE)", "in", ""),
                                 ("Where the finished labels appear  (Generated_Labels)", "out", ""),
                                 ("Where the profiles are kept  (profiles)", "pro", "")], None)
    if r:
        path = {"in": INPUT_FOLDER, "out": OUTPUT_FOLDER, "pro": PROFILES_DIR}[r]
        os.makedirs(path, exist_ok=True)
        open_path(path)


def make_labels_now(profiles):
    files = find_data_files()
    if not files:
        notice("warn", "No spreadsheets yet", ["Drop your .xlsx or .csv files into the EXCEL_FILES_HERE folder, then try again.",
                                               "(Use “Open my folders” on the main menu to jump straight there.)"])
        return
    items = [dict(label=p["profile_name"]) for p in profiles]
    r = menu("Make labels with which profile?", items, esc="Cancel")
    if r is None:
        return
    tui_stop()
    print()
    say("ok", "Found %d spreadsheet%s" % (len(files), "" if len(files) == 1 else "s"))
    stats = run_generation(profiles[r], files)
    if stats:
        open_path(stats["path"])
    print()
    try:
        input("  Press Enter to go back to the menu...")
    except EOFError:
        pass
    tui_start()


def cmd_configure(args):
    if not IS_TTY:
        print("Please run the configure script from a normal terminal / command window.")
        return 1
    tui_start()
    try:
        while True:
            profiles = load_profiles()
            items = [dict(label="Create a new profile", value="new",
                          desc="Step-by-step: teach the app about a new entry form.  Best for a new exhibition or a different school."),
                     dict(label="Edit a profile", value="edit", tag="%d saved" % len(profiles),
                          desc="Change which columns are used, the look of the labels, or rename / copy / delete a profile."),
                     dict(label="Make labels now", value="make",
                          desc="Builds the PDF from the spreadsheets in EXCEL_FILES_HERE."),
                     dict(label="Open my folders", value="folders", desc="Jump to the folders where spreadsheets, profiles and finished labels live."),
                     dict(label="Quit", value="quit")]
            for bad in BAD_PROFILES:
                items.insert(2, dict(label="(unreadable file: %s)" % bad, sep=True))
            r = menu("What would you like to do?", items, esc="Quit",
                     subtitle="Use the ↑ ↓ arrow keys and press Enter.  Nothing you do here can break your spreadsheets.")
            if r is None or items[r].get("value") == "quit":
                break
            v = items[r]["value"]
            if v == "new":
                wizard_new()
            elif v == "edit":
                if not profiles:
                    notice("info", "No profiles yet", ["Choose “Create a new profile” first."])
                    continue
                q = menu("Which profile?", [dict(label=p["profile_name"], tag="%d columns" % len([x for x in p["fields"].values() if x])) for p in profiles],
                         crumb="Edit a profile", esc="Back")
                if q is not None:
                    profile_hub(profiles[q])
            elif v == "make":
                if not profiles:
                    notice("info", "No profiles yet", ["Choose “Create a new profile” first."])
                    continue
                make_labels_now(profiles)
            elif v == "folders":
                folders_menu()
    finally:
        tui_stop()
    print()
    print("\n".join(banner()))
    print("\n  " + sty("All done.", GOOD) + "  Run the labels script whenever your spreadsheets are ready.\n")
    return 0


# ════════════════════════════════════════════════════════════════════════════
def main():
    ap = argparse.ArgumentParser(description="Exhibition Label Studio")
    ap.add_argument("mode", nargs="?", default="generate", choices=["generate", "configure"],
                    help="generate (default) makes the labels; configure opens the set-up screen")
    ap.add_argument("--profile", help="use this profile name instead of choosing")
    ap.add_argument("--limit", type=int, help="only make the first N labels (for testing)")
    args = ap.parse_args()
    try:
        code = cmd_configure(args) if args.mode == "configure" else cmd_generate(args)
    except KeyboardInterrupt:
        tui_stop()
        print("\n  Cancelled.\n")
        code = 1
    sys.exit(code)


if __name__ == "__main__":
    main()
